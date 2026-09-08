"""Tests that check behaviour, not that the code ran.

Each one asserts something a human would actually care about: the right value
came out of the real PDF, a bad form gets rejected for the right reason, and the
workbook still has its colours after we write to it.
"""
import os
import re
import shutil
import sys
import zipfile
from datetime import date

import openpyxl
import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import config
import epif_parser
import log_writer
import validators

SAMPLES = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "samples")
FILLED = os.path.join(SAMPLES, "Prusa_EPIF__2799_PG000025831.pdf")
BLANK = os.path.join(SAMPLES, "EPIF_TEMPLATE_blank.pdf")
WORKBOOK = os.path.join(SAMPLES, "Purchasing-Log.xlsx")


@pytest.fixture
def filled():
    with open(FILLED, "rb") as handle:
        return epif_parser.parse_epif(handle.read())


@pytest.fixture
def workbook(tmp_path):
    copy = tmp_path / "Purchasing-Log.xlsx"
    shutil.copy(WORKBOOK, copy)
    return str(copy)


# --- parsing ------------------------------------------------------------------

def test_reads_the_real_values_off_the_prusa_epif(filled):
    assert filled["item_description"] == "Prusa CORE One L+ INDX 4-Tool"
    assert filled["vendor"] == "Prusa"
    assert filled["vendor_contact_email"] == "info@prusa3d.com"
    assert filled["project_id"] == "PG000025831"
    assert filled["fund"] == "150"
    assert filled["delivery_room"] == "ERB 212"


def test_dollar_sign_amount_becomes_a_number(filled):
    assert filled["total_price"] == 2799.0


def test_two_digit_date_is_understood(filled):
    assert filled["date_of_purchase"] == date(2026, 9, 3)


def test_link_is_pulled_out_of_the_purpose_paragraph(filled):
    assert filled["link"].startswith("https://www.prusa3d.com/")


def test_ticked_checkbox_maps_to_the_dropdown_wording(filled):
    assert filled["category"] == "Research/Lab Supplies (3105)"
    assert filled["category_error"] is None


def test_blank_template_reports_no_category_ticked():
    with open(BLANK, "rb") as handle:
        parsed = epif_parser.parse_epif(handle.read())
    assert parsed["category"] is None
    assert "no EPIF category" in parsed["category_error"]


def test_a_pdf_with_no_form_fields_is_rejected_loudly(tmp_path):
    from pypdf import PdfReader, PdfWriter
    flat = tmp_path / "flat.pdf"
    writer = PdfWriter()
    writer.add_page(PdfReader(FILLED).pages[0])
    # drop the AcroForm, imitating a print-to-PDF
    writer._root_object.pop("/AcroForm", None)
    with open(flat, "wb") as handle:
        writer.write(handle)
    with pytest.raises(epif_parser.FlattenedPdfError):
        epif_parser.parse_epif(flat.read_bytes())


# --- validation ---------------------------------------------------------------

def test_prusa_epif_is_rejected_because_no_payment_box_is_ticked(filled):
    problems = validators.validate(filled, requester_name="Isaac")
    assert any("P-card" in p for p in problems)


def test_unknown_fund_is_named_in_the_complaint(filled):
    filled["fund"] = "999"
    problems = validators.validate(filled, requester_name="Isaac")
    assert any("999" in p for p in problems)


def test_unmapped_slack_user_is_its_own_problem(filled):
    problems = validators.validate(filled, requester_name=None)
    assert any("Slack ID" in p for p in problems)


def test_a_complete_form_passes(filled):
    filled["payment_method"] = "P-card"
    assert validators.validate(filled, requester_name="Isaac") == []


# --- writing ------------------------------------------------------------------

def _x14_blocks(path):
    with zipfile.ZipFile(path) as archive:
        xml = archive.read(config.SHEET_XML).decode("utf-8")
    return len(re.findall(r"<x14:conditionalFormatting", xml))


def test_the_first_empty_row_is_17(workbook):
    with zipfile.ZipFile(workbook) as archive:
        xml = archive.read(config.SHEET_XML).decode("utf-8")
    assert log_writer.find_first_empty_row(xml) == 17


def test_written_values_land_in_the_right_cells(workbook, filled):
    row = log_writer.append_row(
        log_writer.build_row(filled, "Isaac"), workbook_path=workbook
    )
    sheet = openpyxl.load_workbook(workbook)["Order Log"]
    assert sheet[f"B{row}"].value == "Isaac"
    assert sheet[f"C{row}"].value == "Prusa CORE One L+ INDX 4-Tool"
    assert sheet[f"H{row}"].value == 2799
    assert sheet[f"K{row}"].value == "info@prusa3d.com"
    assert sheet[f"R{row}"].value == "PG000025831"
    assert sheet[f"S{row}"].value == 150


def test_the_date_is_a_real_date_not_a_string(workbook, filled):
    row = log_writer.append_row(
        log_writer.build_row(filled, "Isaac"), workbook_path=workbook
    )
    sheet = openpyxl.load_workbook(workbook)["Order Log"]
    assert sheet[f"M{row}"].value.date() == date(2026, 9, 3)


def test_status_colouring_survives_the_write(workbook, filled):
    before = _x14_blocks(workbook)
    assert before > 0, "fixture should start with the colour rules intact"
    log_writer.append_row(
        log_writer.build_row(filled, "Isaac"), workbook_path=workbook
    )
    assert _x14_blocks(workbook) == before


def test_openpyxl_would_have_destroyed_it(workbook):
    """The reason log_writer does zip surgery. If this ever fails, simplify."""
    openpyxl.load_workbook(workbook).save(workbook)
    assert _x14_blocks(workbook) == 0


def test_order_id_and_status_formulas_are_left_alone(workbook, filled):
    row = log_writer.append_row(
        log_writer.build_row(filled, "Isaac"), workbook_path=workbook
    )
    sheet = openpyxl.load_workbook(workbook)["Order Log"]
    assert sheet[f"Z{row}"].value.startswith("=IF(COUNTA(")


def test_writing_a_formula_column_is_refused(workbook):
    with zipfile.ZipFile(workbook) as archive:
        xml = archive.read(config.SHEET_XML).decode("utf-8")
    with pytest.raises(ValueError):
        log_writer.apply_row(xml, 17, {"A": 99})


def test_two_orders_go_to_consecutive_rows(workbook, filled):
    first = log_writer.append_row(log_writer.build_row(filled, "Isaac"), workbook)
    second = log_writer.append_row(log_writer.build_row(filled, "Smeet"), workbook)
    assert second == first + 1


def test_open_in_excel_blocks_the_write(workbook, filled):
    lock = os.path.join(os.path.dirname(workbook), "~$Purchasing-Log.xlsx")
    open(lock, "w").close()
    with pytest.raises(log_writer.WorkbookLockedError):
        log_writer.append_row(log_writer.build_row(filled, "Isaac"), workbook)
