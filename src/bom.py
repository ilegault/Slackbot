"""Line items parsing, formatting, validation, and BOM workbook generation.

WHY THIS EXISTS:
----------------
EPIFs have exactly one 'What is being purchased' description and one 'Amount of Purchase'.
When a lab member orders multiple parts from one vendor (e.g. five shaft couplings from
Ruland), sending purchasing five bare links in an email forces the buyer and Tina to
reconstruct quantities, unit costs, and part numbers by hand.

This module provides the pure domain logic for multi-item orders:
1. Parsing pasted line items (pipe-separated or tab-separated straight from Excel/Sheets)
   with precise per-line error reporting.
2. Formatting line items back to standard text for pre-filling edit modals.
3. Checking line item totals against request total ($0.01 tolerance).
4. Building an in-memory openpyxl workbook (BOM). Formulas are used to match the lab's hand-made BOM
   so purchasing can adjust a quantity. (Accepted cost: Slack's preview shows formula cells blank
   because openpyxl stores no cached results). The sheet has no shipping row, so its Total can be
   below the EPIF amount by the shipping entered, and that this is intended.
5. Hyperlinking URL links in openpyxl on the Item cell so Tina can click straight through to parts.
6. Generating standardized file names (NNNN_<Vendor>_BOM.xlsx) with safe slugging.
7. Computing concise change descriptions for card edit threads.

All functions in this module are strictly pure (no Slack API, no filesystem I/O)
to guarantee isolation and fast, deterministic testability.
"""
import io
import re
from typing import Any, Dict, List, Optional, Tuple

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

try:
    from . import epif_parser
except ImportError:
    import epif_parser


def parse_line_items(text: str) -> Tuple[List[Dict[str, Any]], float, List[str]]:
    """Parse pasted line items text into a list of item dicts, shipping amount, and errors.

    Input format: one item per line, separated by '|' or tab '\\t':
      qty | name | part_number | unit_price | link | description
    Optional shipping line:
      shipping | <amount>

    Returns:
      (items, shipping, errors)
    """
    items: List[Dict[str, Any]] = []
    shipping = 0.0
    shipping_seen = False
    errors: List[str] = []

    if not text:
        return items, shipping, errors

    lines = text.splitlines()
    for idx, raw_line in enumerate(lines):
        line_no = idx + 1
        line = raw_line.strip()
        if not line:
            continue

        if "|" in line:
            fields = [f.strip() for f in line.split("|")]
        else:
            fields = [f.strip() for f in line.split("\t")]

        # Check for shipping line
        if fields and fields[0].lower() == "shipping":
            if shipping_seen:
                errors.append(f"line {line_no}: second shipping line is not allowed")
            elif len(fields) < 2:
                errors.append(f"line {line_no}: shipping line must include an amount")
            elif len(fields) > 2:
                errors.append(f"line {line_no}: shipping line has too many fields")
            else:
                amt_str = fields[1]
                p = epif_parser.parse_money(amt_str)
                if p is None:
                    errors.append(f'line {line_no}: shipping amount "{amt_str}" is not a valid number')
                elif p < 0:
                    errors.append(f'line {line_no}: shipping amount "{amt_str}" must be ≥ 0')
                else:
                    shipping = p
                    shipping_seen = True
            continue

        # Item line: 4 to 6 fields
        if len(fields) < 4:
            errors.append(
                f"line {line_no}: expected at least 4 fields (qty | name | part # | unit price), got {len(fields)}"
            )
            continue
        if len(fields) > 6:
            errors.append(
                f"line {line_no}: too many fields (expected up to 6, got {len(fields)})"
            )
            continue

        qty_str = fields[0]
        name = fields[1]
        part_number = fields[2]
        price_str = fields[3]
        link = fields[4] if len(fields) > 4 else ""
        description = fields[5] if len(fields) > 5 else ""

        line_valid = True

        # Validate qty (positive whole number)
        try:
            qty_val = int(qty_str)
            if qty_val <= 0:
                errors.append(f'line {line_no}: quantity "{qty_str}" must be a positive whole number')
                line_valid = False
        except (ValueError, TypeError):
            errors.append(f'line {line_no}: quantity "{qty_str}" must be a positive whole number')
            line_valid = False
            qty_val = 0

        # Validate name (non-empty)
        if not name:
            errors.append(f"line {line_no}: item name cannot be empty")
            line_valid = False

        # Validate unit_price (number >= 0)
        p = epif_parser.parse_money(price_str)
        if p is None:
            errors.append(f'line {line_no}: unit price "{price_str}" is not a number')
            line_valid = False
            price_val = 0.0
        elif p < 0:
            errors.append(f'line {line_no}: unit price "{price_str}" must be ≥ 0')
            line_valid = False
            price_val = 0.0
        else:
            price_val = p

        if line_valid:
            items.append({
                "qty": qty_val,
                "name": name,
                "part_number": part_number,
                "unit_price": price_val,
                "link": link,
                "description": description,
            })

    if len(items) > 25:
        errors.append(f"too many item lines ({len(items)} items, maximum is 25)")

    return items, shipping, errors


def format_line_items(items: List[Dict[str, Any]], shipping: float = 0.0) -> str:
    """Format line items and shipping back to pipe-separated text for pre-filling modals."""
    lines: List[str] = []
    for item in items:
        qty = item.get("qty", 1)
        name = item.get("name", "")
        part = item.get("part_number", "")
        price = float(item.get("unit_price", 0.0))
        link = item.get("link", "")
        desc = item.get("description", "")

        if desc:
            line = f"{qty} | {name} | {part} | {price:.2f} | {link} | {desc}"
        elif link:
            line = f"{qty} | {name} | {part} | {price:.2f} | {link}"
        else:
            line = f"{qty} | {name} | {part} | {price:.2f}"
        lines.append(line)

    if shipping > 0.0:
        lines.append(f"shipping | {shipping:.2f}")

    return "\n".join(lines)


def check_total(items: List[Dict[str, Any]], shipping: float, total_price: Any) -> Optional[str]:
    """Check if sum(qty * unit_price) + shipping matches total_price within $0.01.

    Returns None if within tolerance, or an error sentence showing both numbers.
    """
    if total_price is None:
        target = 0.0
    elif isinstance(total_price, (int, float)):
        target = float(total_price)
    else:
        parsed = epif_parser.parse_money(str(total_price))
        target = parsed if parsed is not None else 0.0

    calculated = sum(int(item["qty"]) * float(item["unit_price"]) for item in items) + float(shipping or 0.0)
    if abs(calculated - target) <= 0.01 + 1e-9:
        return None
    return (
        f"Line items total (${calculated:.2f}) does not match request total (${target:.2f})."
    )


def needs_bom(items: List[Dict[str, Any]]) -> bool:
    """Return True if request has two or more line items."""
    return len(items) >= 2


def describe_changes(old_request: Dict[str, Any], new_request: Dict[str, Any]) -> List[str]:
    """Return concise change fragments between old_request and new_request."""
    fragments: List[str] = []

    # 1. Total price
    def _parse_p(val: Any) -> Optional[float]:
        if val is None:
            return None
        if isinstance(val, (int, float)):
            return float(val)
        return epif_parser.parse_money(str(val))

    old_p = _parse_p(old_request.get("total_price", old_request.get("Amount of Purchase")))
    new_p = _parse_p(new_request.get("total_price", new_request.get("Amount of Purchase")))
    if old_p is not None and new_p is not None and abs(old_p - new_p) > 0.005:
        fragments.append(f"Total ${old_p:.2f} → ${new_p:.2f}")
    elif old_p is not None and new_p is None:
        fragments.append(f"Total ${old_p:.2f} removed")
    elif old_p is None and new_p is not None:
        fragments.append(f"Total set to ${new_p:.2f}")

    # 2. Line items & shipping
    old_items = old_request.get("items") or []
    new_items = new_request.get("items") or []
    old_ship = float(old_request.get("shipping") or 0.0)
    new_ship = float(new_request.get("shipping") or 0.0)

    if len(old_items) != len(new_items):
        fragments.append(f"items {len(old_items)} → {len(new_items)}")
    elif old_items != new_items or abs(old_ship - new_ship) > 0.005:
        fragments.append("items changed")

    # 3. Scalar text fields
    scalar_fields = [
        (["item_description", "What is being purchased"], "item description"),
        (["purpose", "Purpose"], "purpose"),
        (["link", "Link"], "link"),
        (["project_id", "Project ID Number"], "project ID"),
        (["fund", "Fund"], "fund"),
        (["category"], "category"),
        (["delivery_room", "room_address", "room address"], "delivery room"),
        (["vendor_contact_name", "Vendor Name"], "vendor contact name"),
        (["vendor_contact_email", "Email add"], "vendor contact email"),
        (["payment_method", "how_buying", "How Buying"], "payment method"),
        (["asset_id", "Asset ID"], "asset ID"),
        (["name_of_system", "Name of System"], "name of system"),
    ]

    for aliases, label in scalar_fields:
        old_val = ""
        for k in aliases:
            if k in old_request and old_request[k] is not None:
                old_val = str(old_request[k]).strip()
                break

        new_val = ""
        for k in aliases:
            if k in new_request and new_request[k] is not None:
                new_val = str(new_request[k]).strip()
                break

        if old_val != new_val:
            if old_val and new_val:
                fragments.append(f"{label} '{old_val}' → '{new_val}'")
            elif new_val:
                fragments.append(f"{label} set to '{new_val}'")
            else:
                fragments.append(f"{label} cleared")

    return fragments


def bom_filename(row: Optional[int], vendor: str) -> str:
    """Generate standardized BOM filename.

    Draft: DRAFT_{slug}_BOM.xlsx
    Approved: NNNN_{slug}_BOM.xlsx (zero-padded to 4 digits)
    """
    v_clean = (vendor or "").strip()
    v_slug = re.sub(r"\s+", "-", v_clean)
    v_slug = re.sub(r"[^A-Za-z0-9\-]", "", v_slug)
    v_slug = re.sub(r"-+", "-", v_slug).strip("-")
    v_slug = v_slug[:40].rstrip("-")
    if not v_slug:
        v_slug = "Vendor"

    if row is None:
        return f"DRAFT_{v_slug}_BOM.xlsx"
    return f"{int(row):04d}_{v_slug}_BOM.xlsx"


def build_bom_workbook(
    request: Dict[str, Any],
    items: List[Dict[str, Any]],
) -> bytes:
    """Build an in-memory BOM spreadsheet with openpyxl and return its bytes.

    Layout matches the lab's hand-made BOM.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "BOM"

    # Font definitions
    aptos_regular = Font(name="Aptos Narrow", size=12)
    aptos_bold = Font(name="Aptos Narrow", size=12, bold=True)
    aptos_white_bold = Font(name="Aptos Narrow", size=12, bold=True, color="FFFFFF")
    link_font = Font(name="Aptos Narrow", size=12, color="467886", underline="single")

    # Fill definitions
    fill_0B3041 = PatternFill(start_color="0B3041", end_color="0B3041", fill_type="solid")
    fill_DCEAF7 = PatternFill(start_color="DCEAF7", end_color="DCEAF7", fill_type="solid")
    fill_104862 = PatternFill(start_color="104862", end_color="104862", fill_type="solid")

    # Header block
    vendor = str(request.get("vendor") or "").strip()
    if not vendor:
        vendor = "Unknown Vendor"

    ws.merge_cells("A1:B1")
    ws["A1"] = "BILL OF MATERIALS"
    ws["A1"].font = aptos_white_bold
    ws["A1"].fill = fill_0B3041

    ws["A2"] = "Total Cost"
    ws["A2"].font = aptos_bold
    ws["A2"].fill = fill_DCEAF7

    total_row = 5 + len(items)
    ws["B2"] = f"=G{total_row}"
    ws["B2"].font = aptos_regular
    ws["B2"].fill = fill_DCEAF7
    ws["B2"].number_format = '"$"#,##0.00'
    ws["B2"].alignment = Alignment(horizontal="left")

    # Table Header Row (Row 4)
    headers = ["Item", "Description", "Product #", "Vendor", "Unit Cost", "Quantity", "Total Cost", "Notes"]
    for col_idx, col_name in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=col_idx, value=col_name)
        cell.font = aptos_white_bold
        cell.fill = fill_104862

    current_row = 5
    for item in items:
        qty = int(item.get("qty", 1))
        name = str(item.get("name", "")).strip()
        part = str(item.get("part_number", "")).strip()
        desc = str(item.get("description", "")).strip()
        link_str = str(item.get("link", "")).strip()
        unit_cost = float(item.get("unit_price", 0.0))

        # Col 1: Item
        c1 = ws.cell(row=current_row, column=1, value=name)
        if link_str.startswith("http://") or link_str.startswith("https://"):
            c1.hyperlink = link_str
            c1.font = link_font
        else:
            c1.font = aptos_regular

        # Col 2: Description
        c2 = ws.cell(row=current_row, column=2, value=desc)
        c2.font = aptos_regular

        # Col 3: Product #
        c3 = ws.cell(row=current_row, column=3, value=part)
        c3.font = aptos_regular

        # Col 4: Vendor
        c4 = ws.cell(row=current_row, column=4, value=vendor)
        c4.font = aptos_regular

        # Col 5: Unit Cost
        c5 = ws.cell(row=current_row, column=5, value=unit_cost)
        c5.font = aptos_regular
        c5.number_format = '"$"#,##0.00'

        # Col 6: Quantity
        c6 = ws.cell(row=current_row, column=6, value=qty)
        c6.font = aptos_regular

        # Col 7: Total Cost
        c7 = ws.cell(row=current_row, column=7, value=f"=E{current_row}*F{current_row}")
        c7.font = aptos_regular
        c7.number_format = '"$"#,##0.00'

        # Col 8: Notes (Empty)
        ws.cell(row=current_row, column=8, value=None)

        current_row += 1

    # Total row
    for col_idx in range(1, 9):
        c = ws.cell(row=total_row, column=col_idx)
        c.fill = fill_104862
        if col_idx == 6:
            c.value = "Total"
            c.font = aptos_white_bold
        elif col_idx == 7:
            c.value = f"=SUM(G5:G{total_row-1})"
            c.font = aptos_white_bold
            c.number_format = '"$"#,##0.00'

    # Freeze panes
    ws.freeze_panes = "A5"

    # Column widths
    col_widths = {
        "A": 35,
        "B": 68,
        "C": 14,
        "D": 17,
        "E": 12,
        "F": 10,
        "G": 13,
        "H": 47,
    }
    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = width

    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()
