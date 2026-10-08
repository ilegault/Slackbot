"""Tests for Ticket 101: @Purchasing approved approves a posted card in the thread.

WHY THIS EXISTS:
----------------
Per ADR 0015 (Decisions 1 and 4):
The keyword and the button must approve the same thing. Today @Purchasing approved
reaches lifecycle.handle_epif_processing with no card_ts and no posted_payload.
The function previously searched the thread for a PDF first, so a re-posted quote
in the thread derailed the keyword by trying to parse it as an EPIF form.
Now, when neither card_ts, posted_payload, nor direct_file is given, handle_epif_processing
calls slack_io.find_card_in_thread first. If a card in the 'posted' state is found,
card_ts and posted_payload are populated from it, taking the button path and skipping
the thread PDF search.
"""
import datetime
import io
import json
import os
from unittest.mock import MagicMock

import openpyxl
import pypdf
import pytest

from src import blocks, config, epif_filler, epif_parser, lifecycle, queue_worker, roster, slack_io

FIXTURE_PATH = os.path.join(
    os.path.dirname(__file__), "fixtures", "EPIF_TEMPLATE_HIRST.pdf"
)


@pytest.fixture(autouse=True)
def clean_roster(tmp_path, monkeypatch):
    """Ensure an isolated roster for testing."""
    roster_file = str(tmp_path / "roster.json")
    monkeypatch.setattr(roster, "ROSTER_PATH", roster_file)
    roster._DATA = None
    data = {
        "admins": ["U_ADMIN"],
        "approvers": ["U_CHARLIE"],
        "buyers": ["U_BUYER"],
        "requesters": {
            "U_ADMIN": "Isaac",
            "U_CHARLIE": "Charlie H.",
            "U_BUYER": "Dylan",
            "U_REQ": "Alex",
        },
        "vendors": ["Thorlabs", "MKS Instruments"],
    }
    with open(roster_file, "w", encoding="utf-8") as f:
        json.dump(data, f)
    return roster_file


@pytest.fixture
def temp_workbook(tmp_path, monkeypatch):
    """Create a temporary test workbook."""
    copy_path = str(tmp_path / "Purchasing-Log.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Purchasing Log"
    for _ in range(1, 35):
        ws.append([""] * 30)
    wb.save(copy_path)
    monkeypatch.setattr(config, "WORKBOOK_PATH", copy_path)
    return copy_path


@pytest.fixture
def sync_queue(monkeypatch):
    """Synchronously drain queue_worker submit_write_task calls."""
    def _sync(action_fn, channel="", thread_ts="", user_id="", task_type="append",
              description="", success_callback=None, failure_callback=None, client=None):
        try:
            res = action_fn()
        except Exception as exc:
            if failure_callback:
                failure_callback(exc)
            return
        if success_callback:
            success_callback(res)
    monkeypatch.setattr(queue_worker, "submit_write_task", _sync)


@pytest.fixture
def dirs(tmp_path, monkeypatch):
    """Set up temporary storage directories."""
    out = {}
    for name in ("BOMS_DIR", "QUOTES_DIR", "EPIFS_DIR"):
        d = str(tmp_path / name)
        os.makedirs(d)
        monkeypatch.setattr(config, name, d)
        out[name] = d
    monkeypatch.setattr(config, "EPIF_TEMPLATE_PATH", FIXTURE_PATH)
    monkeypatch.setattr(config, "ADMIN_ALERT_CHANNEL", "C_ADMIN")
    return out


def test_keyword_approves_posted_card_with_quote_in_thread(temp_workbook, dirs, sync_queue, monkeypatch):
    """AC 1 & 2: Thread holds a /new-purchase card in 'posted' state and a bot-posted quote PDF.

    - Calling handle_epif_processing as the mention handler does writes a row with the
      card's item, vendor, and total.
    - No 'Error processing' text is posted.
    - The card is updated in place to the 'approved' state (chat_update called with ts = card_ts).
    - No second card was posted with chat_postMessage + blocks.
    """
    # 1. Real form-less PDF bytes built using pypdf.PdfWriter
    writer = pypdf.PdfWriter()
    writer.add_blank_page(width=612, height=792)
    buf = io.BytesIO()
    writer.write(buf)
    quote_pdf_bytes = buf.getvalue()

    # 2. Spy on epif_parser.read_fields and slack_io.find_epif_in_thread
    real_read_fields = epif_parser.read_fields
    read_fields_spy = MagicMock(side_effect=real_read_fields)
    monkeypatch.setattr(epif_parser, "read_fields", read_fields_spy)

    find_epif_spy = MagicMock(side_effect=slack_io.find_epif_in_thread)
    monkeypatch.setattr(slack_io, "find_epif_in_thread", find_epif_spy)

    # 3. Card data from /new-purchase
    quote_filename = "Quote_Thorlabs_12345.pdf"
    file_id = "F_QUOTE_1"
    card_ts = "100.2"
    thread_ts = "100.0"

    req_data = {
        "parsed": {
            "item_description": "Precision Optical Rail 500mm",
            "total_price": 349.50,
            "vendor": "Thorlabs",
            "category": "Research/Lab Supplies (3105)",
            "category_error": None,
            "project_id": "PG000025831",
            "fund": "133",
            "delivery_room": "ERB 212",
            "purpose": "Beam steering setup",
            "payment_method": "Workday",
            "date_of_purchase": "2026-09-17",
            "link": "https://thorlabs.com/item123",
            "vendor_contact_email": "orders@thorlabs.com",
            "vendor_contact_name": "Thorlabs Rep",
            "name_of_system": "",
            "asset_id": "",
        },
        "requester": "Alex",
        "user_id": "U_REQ",
        "is_pending_name": False,
        "source": "modal",
        "thread_ts": thread_ts,
        "attachments": [
            {
                "id": file_id,
                "name": quote_filename,
                "role": "quote",
            }
        ],
    }

    card_blocks = blocks.build_request_blocks("posted", req_data)

    card_message = {
        "ts": card_ts,
        "user": "B_BOT",
        "text": "New Purchase Request from Alex",
        "metadata": {
            "event_type": "purchase_request",
            "event_payload": req_data,
        },
        "blocks": card_blocks,
    }

    bot_quote_message = {
        "ts": "100.3",
        "user": "B_BOT",
        "text": f"📎 Attached Quote: {quote_filename}",
        "files": [
            {
                "id": file_id,
                "name": quote_filename,
                "mimetype": "application/pdf",
                "url_private_download": "https://slack.com/fake/download",
            }
        ],
    }

    parent_message = {
        "ts": thread_ts,
        "user": "U_REQ",
        "text": "I'd like to order an optical rail.",
    }

    # 4. Fake Slack client
    client = MagicMock()
    posted_messages = []
    chat_updates = []

    def fake_post_message(**kwargs):
        posted_messages.append(kwargs)
        return {"ok": True, "ts": "100.99"}

    client.chat_postMessage.side_effect = fake_post_message

    def fake_chat_update(**kwargs):
        chat_updates.append(kwargs)
        return {"ok": True, "ts": kwargs.get("ts")}

    client.chat_update.side_effect = fake_chat_update

    def fake_replies(channel, ts, **kwargs):
        return {
            "ok": True,
            "messages": [parent_message, card_message, bot_quote_message],
        }

    client.conversations_replies.side_effect = fake_replies
    client.files_info.return_value = {
        "ok": True,
        "file": {
            "id": file_id,
            "name": quote_filename,
            "url_private_download": "https://slack.com/fake/download",
        },
    }
    client.conversations_open.return_value = {"ok": True, "channel": {"id": "D_BUYER"}}
    client.auth_test.return_value = {"ok": True, "user_id": "B_BOT"}

    monkeypatch.setattr(slack_io, "download_file", lambda f: quote_pdf_bytes)
    monkeypatch.setattr(slack_io, "download", lambda f: quote_pdf_bytes)

    say = MagicMock(side_effect=lambda text, thread_ts=None, **kw: fake_post_message(text=text, thread_ts=thread_ts, **kw))

    # 5. Call handle_epif_processing as the mention handler does (@Purchasing approved)
    lifecycle.handle_epif_processing(
        client=client,
        say=say,
        channel="C_PURCHASING",
        thread_ts=thread_ts,
        approver="U_CHARLIE",
        event_ts="100.4",
        direct_file=None,
        direct_poster=None,
        posted_payload=None,
        card_ts=None,
    )

    # 6. Assertions for AC 1:
    # Row is written with the card's item, vendor and total
    wb = openpyxl.load_workbook(temp_workbook)
    ws = wb.active
    logged_row = None
    for r in range(2, 50):
        if ws[f"C{r}"].value == "Precision Optical Rail 500mm":
            logged_row = r
            break

    assert logged_row is not None, "Expected row for 'Precision Optical Rail 500mm' was not written to workbook!"
    assert ws[f"I{logged_row}"].value == "Thorlabs"
    assert float(ws[f"{config.COLUMN_TOTAL_PRICE}{logged_row}"].value) == 349.50

    # No 'Error processing' text is posted
    for msg in posted_messages:
        text = msg.get("text", "")
        assert not text.startswith("Error processing"), f"Unexpected error message: {text}"
        assert "Error processing" not in text, f"Unexpected error message: {text}"

    # Assert read_fields was never called on the quote, and find_epif_in_thread was not called
    assert read_fields_spy.call_count == 0, "epif_parser.read_fields must not be called on the quote!"
    assert find_epif_spy.call_count == 0, "slack_io.find_epif_in_thread must not be called when a posted card exists!"

    # 7. Assertions for AC 2:
    # Card is updated in place to the 'approved' state
    assert len(chat_updates) >= 1, "chat_update should have been called to update the card!"
    assert chat_updates[0]["ts"] == card_ts, f"chat_update was called with ts={chat_updates[0]['ts']}, expected {card_ts}"

    # Assert no second card was posted with chat_postMessage + blocks
    for msg in posted_messages:
        assert "blocks" not in msg, f"Found chat_postMessage with blocks: {msg}"


def test_keyword_approves_posted_card_when_epif_also_in_thread(temp_workbook, dirs, sync_queue, monkeypatch):
    """AC 3: A thread with a posted card AND a person's dropped EPIF approves the card.

    Assert the row's vendor is the card's, not the EPIF's.
    """
    with open(FIXTURE_PATH, "rb") as f:
        template_bytes = f.read()

    epif_data = {
        "item_description": "Vacuum Flange Assembly",
        "purpose": "Chamber test",
        "total_price": 850.00,
        "vendor": "MKS Instruments",
        "vendor_contact_name": "MKS Rep",
        "vendor_contact_email": "rep@mksinst.com",
        "category": "Research/Lab Supplies (3105)",
        "payment_method": "P-card",
        "date_of_purchase": datetime.date(2026, 9, 23),
        "delivery_room": "ERB 212",
        "project_id": "PG000025831",
        "fund": "133",
        "name_of_system": "",
        "asset_id": "",
    }
    dropped_epif_bytes = epif_filler.fill_epif(template_bytes, epif_data)

    card_ts = "100.2"
    thread_ts = "100.0"

    card_req_data = {
        "parsed": {
            "item_description": "Precision Optical Rail 500mm",
            "total_price": 349.50,
            "vendor": "Thorlabs",
            "category": "Research/Lab Supplies (3105)",
            "category_error": None,
            "project_id": "PG000025831",
            "fund": "133",
            "delivery_room": "ERB 212",
            "purpose": "Beam steering setup",
            "payment_method": "Workday",
            "date_of_purchase": "2026-09-17",
            "link": "https://thorlabs.com/item123",
            "vendor_contact_email": "orders@thorlabs.com",
            "vendor_contact_name": "Thorlabs Rep",
            "name_of_system": "",
            "asset_id": "",
        },
        "requester": "Alex",
        "user_id": "U_REQ",
        "is_pending_name": False,
        "source": "modal",
        "thread_ts": thread_ts,
    }

    card_blocks = blocks.build_request_blocks("posted", card_req_data)

    parent_message = {
        "ts": thread_ts,
        "user": "U_REQ",
        "text": "Please order this optical rail.",
    }

    card_message = {
        "ts": card_ts,
        "user": "B_BOT",
        "text": "New Purchase Request from Alex",
        "metadata": {
            "event_type": "purchase_request",
            "event_payload": card_req_data,
        },
        "blocks": card_blocks,
    }

    dropped_epif_message = {
        "ts": "100.3",
        "user": "U_REQ",  # Posted by a person, not the bot
        "text": "Here is an EPIF for something else.",
        "files": [
            {
                "id": "F_EPIF",
                "name": "MKS_EPIF.pdf",
                "mimetype": "application/pdf",
                "url_private_download": "https://slack.com/fake/epif.pdf",
            }
        ],
    }

    client = MagicMock()
    posted_messages = []
    chat_updates = []

    def fake_post_message(**kwargs):
        posted_messages.append(kwargs)
        return {"ok": True, "ts": "100.99"}

    client.chat_postMessage.side_effect = fake_post_message

    def fake_chat_update(**kwargs):
        chat_updates.append(kwargs)
        return {"ok": True, "ts": kwargs.get("ts")}

    client.chat_update.side_effect = fake_chat_update

    def fake_replies(channel, ts, **kwargs):
        return {
            "ok": True,
            "messages": [parent_message, card_message, dropped_epif_message],
        }

    client.conversations_replies.side_effect = fake_replies
    client.conversations_open.return_value = {"ok": True, "channel": {"id": "D_BUYER"}}
    client.auth_test.return_value = {"ok": True, "user_id": "B_BOT"}

    monkeypatch.setattr(slack_io, "download", lambda f: dropped_epif_bytes)

    say = MagicMock(side_effect=lambda text, thread_ts=None, **kw: fake_post_message(text=text, thread_ts=thread_ts, **kw))

    # Keyword approval with no card_ts, no posted_payload, no direct_file
    lifecycle.handle_epif_processing(
        client=client,
        say=say,
        channel="C_PURCHASING",
        thread_ts=thread_ts,
        approver="U_CHARLIE",
        event_ts="100.4",
        direct_file=None,
        direct_poster=None,
        posted_payload=None,
        card_ts=None,
    )

    # Assert row is written with the card's vendor (Thorlabs), NOT the EPIF's (MKS Instruments)
    wb = openpyxl.load_workbook(temp_workbook)
    ws = wb.active
    logged_row = None
    for r in range(2, 50):
        if ws[f"C{r}"].value == "Precision Optical Rail 500mm":
            logged_row = r
            break

    assert logged_row is not None, "Expected row for card item 'Precision Optical Rail 500mm' not found!"
    assert ws[f"I{logged_row}"].value == "Thorlabs", f"Expected card vendor 'Thorlabs', got {ws[f'I{logged_row}'].value}"
    assert ws[f"I{logged_row}"].value != "MKS Instruments", "EPIF vendor 'MKS Instruments' was logged instead of card vendor!"

    # Assert card was updated in place
    assert len(chat_updates) >= 1
    assert chat_updates[0]["ts"] == card_ts
