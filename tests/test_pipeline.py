"""Tests that check behaviour, not that the code ran.

Each one asserts something a human would actually care about: the right value
came out of the real PDF, a bad form gets rejected for the right reason, and the
workbook still has its colours after we write to it.
"""
import os
import re
import shutil
import sys
import zipfile
from datetime import date
from unittest.mock import MagicMock

import openpyxl
import pytest

# Determine project root and src directory dynamically
_CURRENT_DIR = os.path.dirname(os.path.abspath(__file__))
PROJECT_ROOT = os.path.dirname(_CURRENT_DIR) if os.path.basename(_CURRENT_DIR) == "tests" else _CURRENT_DIR
SRC_DIR = os.path.join(PROJECT_ROOT, "src")
for p in (PROJECT_ROOT, SRC_DIR):
    if p not in sys.path:
        sys.path.insert(0, p)

try:
    from src import app, config, epif_parser, log_writer, validators
except ImportError:
    import app
    import config
    import epif_parser
    import log_writer
    import validators

SAMPLES = os.path.join(PROJECT_ROOT, "samples")
FILLED = os.path.join(SAMPLES, "Prusa_EPIF__2799_PG000025831.pdf")
BLANK = os.path.join(SAMPLES, "EPIF_TEMPLATE_blank.pdf")
WORKBOOK = os.path.join(SAMPLES, "Purchasing-Log.xlsx")


@pytest.fixture
def filled():
    with open(FILLED, "rb") as handle:
        return epif_parser.parse_epif(handle.read())


@pytest.fixture
def workbook(tmp_path):
    copy = tmp_path / "Purchasing-Log.xlsx"
    shutil.copy(WORKBOOK, copy)
    return str(copy)


# --- parsing ------------------------------------------------------------------

def test_reads_the_real_values_off_the_prusa_epif(filled):
    assert filled["item_description"] == "Prusa CORE One L+ INDX 4-Tool"
    assert filled["vendor"] == "Prusa"
    assert filled["vendor_contact_email"] == "info@prusa3d.com"
    assert filled["project_id"] == "PG000025831"
    assert filled["fund"] == "150"
    assert filled["delivery_room"] == "ERB 212"


def test_dollar_sign_amount_becomes_a_number(filled):
    assert filled["total_price"] == 2799.0


def test_two_digit_date_is_understood(filled):
    assert filled["date_of_purchase"] == date(2026, 9, 3)


def test_link_is_pulled_out_of_the_purpose_paragraph(filled):
    assert filled["link"].startswith("https://www.prusa3d.com/")


def test_ticked_checkbox_maps_to_the_dropdown_wording(filled):
    assert filled["category"] == "Research/Lab Supplies (3105)"
    assert filled["category_error"] is None


def test_blank_template_reports_no_category_ticked():
    with open(BLANK, "rb") as handle:
        parsed = epif_parser.parse_epif(handle.read())
    assert parsed["category"] is None
    assert "no EPIF category" in parsed["category_error"]


def test_a_pdf_with_no_form_fields_is_rejected_loudly(tmp_path):
    from pypdf import PdfReader, PdfWriter
    flat = tmp_path / "flat.pdf"
    writer = PdfWriter()
    writer.add_page(PdfReader(FILLED).pages[0])
    # drop the AcroForm, imitating a print-to-PDF
    writer._root_object.pop("/AcroForm", None)
    with open(flat, "wb") as handle:
        writer.write(handle)
    with pytest.raises(epif_parser.FlattenedPdfError):
        epif_parser.parse_epif(flat.read_bytes())


# --- validation ---------------------------------------------------------------

def test_prusa_epif_is_rejected_because_no_payment_box_is_ticked(filled):
    problems = validators.validate(filled, requester_name="Isaac")
    assert any("P-card" in p for p in problems)


def test_unknown_fund_is_named_in_the_complaint(filled):
    filled["fund"] = "999"
    problems = validators.validate(filled, requester_name="Isaac")
    assert any("999" in p for p in problems)


def test_unmapped_slack_user_is_its_own_problem(filled):
    problems = validators.validate(filled, requester_name=None)
    assert any("Slack ID" in p for p in problems)


def test_a_complete_form_passes(filled):
    filled["payment_method"] = "P-card"
    assert validators.validate(filled, requester_name="Isaac") == []


# --- requester resolution and logging -----------------------------------------

def test_resolve_requester_via_explicit_mapping():
    mock_client = MagicMock()
    config.SLACK_USER_TO_REQUESTER["U12345"] = "Isaac"
    try:
        assert app.resolve_requester(mock_client, "U12345") == "Isaac"
    finally:
        config.SLACK_USER_TO_REQUESTER.pop("U12345", None)


def test_resolve_requester_via_users_info_fallback():
    mock_client = MagicMock()
    mock_client.users_info.return_value = {
        "ok": True,
        "user": {
            "name": "isaac.l",
            "profile": {
                "display_name": "Isaac",
                "real_name": "Isaac Legault",
            },
        },
    }
    assert app.resolve_requester(mock_client, "U99999") == "Isaac"


def test_resolve_requester_returns_none_if_unmatched():
    mock_client = MagicMock()
    mock_client.users_info.return_value = {
        "ok": True,
        "user": {
            "name": "random_person",
            "profile": {
                "display_name": "Unknown Person",
                "real_name": "Unknown Person",
            },
        },
    }
    assert app.resolve_requester(mock_client, "U88888") is None


# --- writing ------------------------------------------------------------------

def _x14_blocks(path):
    with zipfile.ZipFile(path) as archive:
        xml = archive.read(config.SHEET_XML).decode("utf-8")
    return len(re.findall(r"<x14:conditionalFormatting", xml))


def test_the_first_empty_row_is_17(workbook):
    with zipfile.ZipFile(workbook) as archive:
        xml = archive.read(config.SHEET_XML).decode("utf-8")
    assert log_writer.find_first_empty_row(xml) == 17


def test_written_values_land_in_the_right_cells(workbook, filled):
    row = log_writer.append_row(
        log_writer.build_row(filled, "Isaac"), workbook_path=workbook
    )
    sheet = openpyxl.load_workbook(workbook)["Order Log"]
    assert sheet[f"B{row}"].value == "Isaac"
    assert sheet[f"C{row}"].value == "Prusa CORE One L+ INDX 4-Tool"
    assert sheet[f"H{row}"].value == 2799
    assert sheet[f"K{row}"].value == "info@prusa3d.com"
    assert sheet[f"R{row}"].value == "PG000025831"
    assert sheet[f"S{row}"].value == 150


def test_the_date_is_a_real_date_not_a_string(workbook, filled):
    row = log_writer.append_row(
        log_writer.build_row(filled, "Isaac"), workbook_path=workbook
    )
    sheet = openpyxl.load_workbook(workbook)["Order Log"]
    assert sheet[f"M{row}"].value.date() == date(2026, 9, 3)


def test_status_colouring_survives_the_write(workbook, filled):
    before = _x14_blocks(workbook)
    assert before > 0, "fixture should start with the colour rules intact"
    log_writer.append_row(
        log_writer.build_row(filled, "Isaac"), workbook_path=workbook
    )
    assert _x14_blocks(workbook) == before


def test_openpyxl_would_have_destroyed_it(workbook):
    """The reason log_writer does zip surgery. If this ever fails, simplify."""
    openpyxl.load_workbook(workbook).save(workbook)
    assert _x14_blocks(workbook) == 0


def test_order_id_and_status_formulas_are_left_alone(workbook, filled):
    row = log_writer.append_row(
        log_writer.build_row(filled, "Isaac"), workbook_path=workbook
    )
    sheet = openpyxl.load_workbook(workbook)["Order Log"]
    assert sheet[f"Z{row}"].value.startswith("=IF(COUNTA(")


def test_writing_a_formula_column_is_refused(workbook):
    with zipfile.ZipFile(workbook) as archive:
        xml = archive.read(config.SHEET_XML).decode("utf-8")
    with pytest.raises(ValueError):
        log_writer.apply_row(xml, 17, {"A": 99})


def test_two_orders_go_to_consecutive_rows(workbook, filled):
    first = log_writer.append_row(log_writer.build_row(filled, "Isaac"), workbook)
    second = log_writer.append_row(log_writer.build_row(filled, "Smeet"), workbook)
    assert second == first + 1


def test_open_in_excel_blocks_the_write(workbook, filled):
    lock = os.path.join(os.path.dirname(workbook), "~$Purchasing-Log.xlsx")
    open(lock, "w").close()
    with pytest.raises(log_writer.WorkbookLockedError):
        log_writer.append_row(log_writer.build_row(filled, "Isaac"), workbook)


# --- epif file saving ---------------------------------------------------------

def test_save_epif_writes_pdf_correctly(tmp_path):
    with open(FILLED, "rb") as f:
        pdf_bytes = f.read()

    save_dir = str(tmp_path / "EPIFs")
    dest = log_writer.save_epif(pdf_bytes, "Sample_EPIF.pdf", target_dir=save_dir)

    assert os.path.exists(dest)
    assert os.path.basename(dest) == "Sample_EPIF.pdf"
    with open(dest, "rb") as f:
        assert f.read() == pdf_bytes


def test_save_epif_handles_missing_extension_and_traversal(tmp_path):
    save_dir = str(tmp_path / "EPIFs")
    dest = log_writer.save_epif(b"%PDF-test", "../../evil_name", target_dir=save_dir)

    assert os.path.exists(dest)
    assert os.path.basename(dest) == "evil_name.pdf"
    assert os.path.dirname(dest) == save_dir


# --- email draft & confirmation flow ------------------------------------------

def test_generate_email_draft_contains_all_key_elements(filled):
    draft = app.generate_email_draft(filled, "Isaac Legault")
    assert "Tina and Ally" in draft
    assert "Prusa" in draft
    assert "$2,799.00" in draft
    assert "PG000025831" in draft
    assert "Fund 150" in draft
    assert "Isaac Legault" in draft
    assert "https://www.prusa3d.com/" in draft


def test_extract_row_from_text():
    assert app.extract_row_from_text("confirmed row 17") == 17
    assert app.extract_row_from_text("confirm #18") == 18
    assert app.extract_row_from_text("package confirmed 19") == 19
    assert app.extract_row_from_text("no row mentioned") is None


def test_extract_price_from_text():
    assert app.extract_price_from_text("Total: $152.49, about to be submitted") == 152.49
    assert app.extract_price_from_text("submitted for $85.00") == 85.0
    assert app.extract_price_from_text("price is 120.50") == 120.50
    assert app.extract_price_from_text("submitted without price") is None


def test_update_row_and_confirm(workbook, filled):
    row = log_writer.append_row(
        log_writer.build_row(filled, "Isaac"), workbook_path=workbook
    )
    today = date(2026, 9, 8)
    log_writer.update_row(row, {config.COLUMN_DATE_CONFIRMED: today}, workbook_path=workbook)

    sheet = openpyxl.load_workbook(workbook)["Order Log"]
    assert sheet[f"V{row}"].value.date() == today

    info = log_writer.get_row_info(row, workbook_path=workbook)
    assert info["requester"] == "Isaac"
    assert info["item_description"] == "Prusa CORE One L+ INDX 4-Tool"


def test_submission_updates_date_processed_and_price(workbook, filled):
    row = log_writer.append_row(
        log_writer.build_row(filled, "Isaac"), workbook_path=workbook
    )
    today = date(2026, 9, 8)
    log_writer.update_row(row, {
        config.COLUMN_DATE_PROCESSED: today,
        config.COLUMN_TOTAL_PRICE: 152.49,
    }, workbook_path=workbook)

    sheet = openpyxl.load_workbook(workbook)["Order Log"]
    assert sheet[f"U{row}"].value.date() == today
    assert sheet[f"H{row}"].value == 152.49


def test_find_latest_unconfirmed_row(workbook, filled):
    row = log_writer.append_row(
        log_writer.build_row(filled, "Isaac"), workbook_path=workbook
    )
    assert log_writer.find_latest_unconfirmed_row_for_requester("Isaac", workbook_path=workbook) == row

    # Now mark confirmed and verify it returns None or older row
    log_writer.update_row(row, {config.COLUMN_DATE_CONFIRMED: date(2026, 9, 8)}, workbook_path=workbook)
    assert log_writer.find_latest_unconfirmed_row_for_requester("Isaac", workbook_path=workbook) is None


# --- confirmation and quote attachments ---------------------------------------

def test_save_confirmation_file(tmp_path):
    conf_bytes = b"%PDF-1.4 Fake Confirmation"
    save_dir = str(tmp_path / "Order-Confirmations")
    dest = log_writer.save_confirmation(conf_bytes, "Keysight_Confirmation.pdf", target_dir=save_dir)

    assert os.path.exists(dest)
    assert os.path.basename(dest) == "Keysight_Confirmation.pdf"
    with open(dest, "rb") as f:
        assert f.read() == conf_bytes


def test_save_quote_file(tmp_path):
    quote_bytes = b"%PDF-1.4 Fake Vendor Quote"
    save_dir = str(tmp_path / "Quotes")
    dest = log_writer.save_quote(quote_bytes, "BH_Cart_Quote.pdf", target_dir=save_dir)

    assert os.path.exists(dest)
    assert os.path.basename(dest) == "BH_Cart_Quote.pdf"
    with open(dest, "rb") as f:
        assert f.read() == quote_bytes


def test_help_message_returns_command_list():
    help_text = app.get_help_message()
    assert "@p-bot approved" in help_text
    assert "@p-bot claim" in help_text
    assert "@p-bot submitted" in help_text
    assert "@p-bot confirmed" in help_text
    assert "@p-bot delivered" in help_text
    assert "@p-bot quote" in help_text


def test_app_home_opened_publishes_view():
    mock_client = MagicMock()
    app.handle_app_home_opened(mock_client, {"user": "U12345"})
    mock_client.views_publish.assert_called_once_with(
        user_id="U12345",
        view=app.APP_HOME_VIEW,
    )
