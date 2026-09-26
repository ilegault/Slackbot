import json
import os
from unittest.mock import MagicMock

import openpyxl
import pytest

from src import app, config, queue_worker, roster, slack_io


@pytest.fixture
def temp_workbook(tmp_path, monkeypatch):
    copy_path = str(tmp_path / "Purchasing-Log.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Purchasing Log"
    for r in range(1, 35):
        ws.append([""] * 30)
    wb.save(copy_path)
    monkeypatch.setattr(config, "WORKBOOK_PATH", copy_path)
    return copy_path

@pytest.fixture
def sync_queue(monkeypatch):
    """Run queue tasks synchronously."""
    def fake_submit(action_fn, channel, thread_ts, user_id, task_type, description, success_callback, failure_callback, client):
        try:
            res = action_fn()
            success_callback(res)
        except Exception as e:
            failure_callback(e)
    monkeypatch.setattr(queue_worker, "submit_write_task", fake_submit)

@pytest.fixture(autouse=True)
def clean_roster(tmp_path, monkeypatch):
    roster_file = str(tmp_path / "roster.json")
    monkeypatch.setattr(roster, "ROSTER_PATH", roster_file)
    roster._DATA = None
    data = {
        "admins": ["U_ADMIN"],
        "approvers": ["U_CHARLIE"],
        "buyers": ["U_BUYER"],
        "requesters": {
            "U_REQ": "Alex"
        },
        "vendors": [{"name": "Grainger", "id": "V1"}],
    }
    with open(roster_file, "w", encoding="utf-8") as f:
        json.dump(data, f)
    roster.load_roster()
    return roster_file

def test_scenario_needs_epif_end_to_end(temp_workbook, sync_queue, monkeypatch, tmp_path):
    monkeypatch.setattr(config, "EPIFS_DIR", str(tmp_path))
    monkeypatch.setattr(config, "PURCHASING_CHANNEL", "C_PURCHASING")
    # We must point EPIF_TEMPLATE_PATH to a valid template in tests/fixtures
    monkeypatch.setattr(config, "EPIF_TEMPLATE_PATH", "tests/fixtures/EPIF_TEMPLATE_HIRST.pdf")

    mock_client = MagicMock()
    mock_ack = MagicMock()
    mock_respond = MagicMock()

    # Mock resolve_requester to match
    monkeypatch.setattr(slack_io, "resolve_requester", lambda c, u: "Alex" if u == "U_REQ" else ("Dylan" if u == "U_BUYER" else None))

    card_ts = "1234.5678"
    thread_ts = "1234.0000"

    val_data = {
        "request": {
            "user_id": "U_REQ",
            "requester": "Alex",
            "assignee_id": "U_BUYER",
            "vendor": "TestVendor",
        },
        "approver": "U_CHARLIE"
    }

    monkeypatch.setattr(slack_io, "find_card_in_thread", lambda *args: (val_data["request"], card_ts, [], "1234.5678"))
    monkeypatch.setattr(slack_io, "get_card_payload", lambda *args: val_data["request"])

    needs_epif_body = {
        "user": {"id": "U_BUYER"},
        "channel": {"id": "C_PURCHASING"},
        "message": {"ts": card_ts},
        "container": {"thread_ts": thread_ts},
        "trigger_id": "trigger_needs_epif",
        "actions": [{"value": json.dumps(val_data)}]
    }

    # Step 1: Assignee presses "This needs an EPIF"
    app.handle_req_needs_epif(mock_ack, needs_epif_body, mock_respond, mock_client)
    mock_ack.assert_called_once()
    mock_client.views_open.assert_called_once()

    open_kwargs = mock_client.views_open.call_args[1]
    view = open_kwargs["view"]
    meta = json.loads(view["private_metadata"])
    assert meta["route"] == "epif"
    assert meta["card_ts"] == card_ts
    assert meta["thread_ts"] == thread_ts

    # Step 2: Assignee fills in Screen 2 (Details) and submits
    mock_ack.reset_mock()
    mock_client.reset_mock()

    stage2_values = {
        "block_item_description": {"item_description": {"value": "Test Item"}},
        "block_total_price": {"total_price": {"value": "100.00"}},
        "block_vendor": {"vendor_choice": {"selected_option": {"value": "None of these — this will be an EPIF order"}}, "vendor_custom": {"value": "TestVendor"}},
        "block_vendor_contact_name": {"vendor_contact_name": {"value": "TestName"}},
        "block_vendor_contact_email": {"vendor_contact_email": {"value": "test@test.com"}},
        "block_date_of_purchase": {"date_of_purchase": {"selected_date": "2026-09-22"}},
        "block_delivery_room": {"delivery_room": {"selected_option": {"value": "ERB 212"}}},

        "block_project_id": {"project_id": {"selected_option": {"value": "PG000025831"}}},
        "block_fund": {"fund": {"selected_option": {"value": "133"}}},
        "block_category": {"category": {"selected_option": {"value": "Supplies"}}},
        "block_purpose": {"purpose": {"value": "Testing Needs EPIF"}},
        "block_payment_method": {"payment_method": {"selected_option": {"value": "EPIF"}}},
    }

    view_submission_body = {
        "user": {"id": "U_BUYER"}
    }
    view_submission = {
        "state": {"values": stage2_values},
        "private_metadata": json.dumps(meta)
    }

    app.handle_stage2_submit(mock_ack, view_submission_body, mock_client, view_submission)

    # Verify _process_interview_completion executed finalizing
    # Should update the card to 'approved'
    update_calls = [c for c in mock_client.chat_update.mock_calls]
    assert len(update_calls) == 1
    update_kwargs = update_calls[0][1] if update_calls[0][1] else update_calls[0].kwargs
    assert update_kwargs["ts"] == card_ts
    assert "Approved" in update_kwargs["text"]

    # And there should be NO chat_postMessage for posting a NEW request card.
    # Actually _say gets called inside with the "Logged to row..." text
    post_calls = [c for c in mock_client.chat_postMessage.mock_calls if ("text" in c[1] and "Logged to row" in c[1]["text"]) or (hasattr(c, "kwargs") and "text" in c.kwargs and "Logged to row" in c.kwargs["text"])]
    assert len(post_calls) == 1

    # Assignee gets a DM with files_upload_v2
    upload_calls = [c for c in mock_client.files_upload_v2.mock_calls]
    assert len(upload_calls) == 1
    upload_kwargs = upload_calls[0][1] if upload_calls[0][1] else upload_calls[0].kwargs
    assert "EPIF" in upload_kwargs["filename"]

    # Verify row was written to temp_workbook
    wb = openpyxl.load_workbook(temp_workbook)
    ws = wb.active
    # First row is 11 for logs based on config, wait config.FIRST_DATA_ROW
    row_num = config.FIRST_DATA_ROW
    row_1 = [cell.value for cell in ws[row_num]]
    assert "Test Item" in row_1

    # Assert a PDF was created in EPIFS_DIR
    files = os.listdir(tmp_path)
    pdf_files = [f for f in files if f.endswith(".pdf")]
    assert len(pdf_files) == 1

def test_needs_epif_missing_template_fails_safely(temp_workbook, sync_queue, monkeypatch, tmp_path):
    monkeypatch.setattr(config, "EPIFS_DIR", str(tmp_path))
    monkeypatch.setattr(config, "PURCHASING_CHANNEL", "C_PURCHASING")
    # Point template to missing file
    monkeypatch.setattr(config, "EPIF_TEMPLATE_PATH", "tests/fixtures/NON_EXISTENT.pdf")

    mock_client = MagicMock()
    mock_ack = MagicMock()

    meta = {
        "channel": "C_PURCHASING",
        "thread_ts": "1234.0000",
        "card_ts": "1234.5678",
        "user_id": "U_REQ",
        "resolved_name": "Alex",
        "approver": "U_CHARLIE",
        "assignee_id": "U_BUYER",
        "route": "epif",
        "vendor_choice": "None of these — this will be an EPIF order",
        "vendor_custom": "TestVendor",
    }

    stage2_values = {
        "block_item_description": {"item_description": {"value": "Test Item"}},
        "block_total_price": {"total_price": {"value": "100.00"}},
        "block_vendor": {"vendor_choice": {"selected_option": {"value": "None of these — this will be an EPIF order"}}, "vendor_custom": {"value": "TestVendor"}},
        "block_vendor_contact_name": {"vendor_contact_name": {"value": "TestName"}},
        "block_vendor_contact_email": {"vendor_contact_email": {"value": "test@test.com"}},
        "block_date_of_purchase": {"date_of_purchase": {"selected_date": "2026-09-22"}},
        "block_delivery_room": {"delivery_room": {"selected_option": {"value": "ERB 212"}}},

        "block_project_id": {"project_id": {"selected_option": {"value": "PG000025831"}}},
        "block_fund": {"fund": {"selected_option": {"value": "133"}}},
        "block_category": {"category": {"selected_option": {"value": "Supplies"}}},
        "block_purpose": {"purpose": {"value": "Testing Missing Template"}},
        "block_payment_method": {"payment_method": {"selected_option": {"value": "EPIF"}}},
    }

    view_submission_body = {"user": {"id": "U_BUYER"}}
    view_submission = {
        "state": {"values": stage2_values},
        "private_metadata": json.dumps(meta)
    }

    app.handle_stage2_submit(mock_ack, view_submission_body, mock_client, view_submission)

    # Should post failure
    post_calls = [c for c in mock_client.chat_postMessage.mock_calls if ("text" in c[1] and "Error" in c[1]["text"]) or (hasattr(c, "kwargs") and "text" in c.kwargs and "Error" in c.kwargs["text"])]
    assert len(post_calls) >= 1

    # Card is NOT updated to approved
    update_calls = [c for c in mock_client.chat_update.mock_calls]
    assert len(update_calls) == 0

    # Row is NOT saved (blanked out)
    wb = openpyxl.load_workbook(temp_workbook)
    ws = wb.active
    row_num = config.FIRST_DATA_ROW
    row_1 = [cell.value for cell in ws[row_num]]
    assert all(not c for c in row_1)
