"""Tests for Ticket 99: The Approve button approves the card it is on.

WHY THIS EXISTS:
----------------
Per ADR 0015 (Decision 1):
A click on Approve reads the request from that card. It never searches the thread
for an EPIF. This fixes the production bug where re-posted vendor quote PDFs in
the thread were mistakenly inspected by find_epif_in_thread before reading the card,
failing with "This PDF has no fillable form fields" and aborting approval.

Per ADR 0012 (Decision 5):
Attached quotes are carried, never read, and archived byte-for-byte upon approval.
"""
import io
import json
import os
from unittest.mock import MagicMock

import openpyxl
import pypdf
import pytest

from src import app, bom, config, epif_parser, queue_worker, roster, slack_io


@pytest.fixture(autouse=True)
def clean_roster(tmp_path, monkeypatch):
    """Ensure an isolated roster for testing."""
    roster_file = str(tmp_path / "roster.json")
    monkeypatch.setattr(roster, "ROSTER_PATH", roster_file)
    roster._DATA = None
    data = {
        "admins": ["U_ADMIN"],
        "approvers": ["U_CHARLIE"],
        "buyers": ["U_BUYER"],
        "requesters": {
            "U_ADMIN": "Isaac",
            "U_CHARLIE": "Charlie H.",
            "U_BUYER": "Dylan",
            "U_REQ": "Alex",
        },
        "vendors": ["Thorlabs"],
    }
    with open(roster_file, "w", encoding="utf-8") as f:
        json.dump(data, f)
    return roster_file


@pytest.fixture
def temp_workbook(tmp_path, monkeypatch):
    """Create a temporary test workbook as in test_40."""
    copy_path = str(tmp_path / "Purchasing-Log.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Purchasing Log"
    for _ in range(1, 35):
        ws.append([""] * 30)
    wb.save(copy_path)
    monkeypatch.setattr(config, "WORKBOOK_PATH", copy_path)
    return copy_path


@pytest.fixture
def sync_queue(monkeypatch):
    """Synchronously drain queue_worker submit_write_task calls."""
    def _sync(action_fn, channel="", thread_ts="", user_id="", task_type="append",
              description="", success_callback=None, failure_callback=None, client=None):
        try:
            res = action_fn()
        except Exception as exc:
            if failure_callback:
                failure_callback(exc)
            return
        if success_callback:
            success_callback(res)
    monkeypatch.setattr(queue_worker, "submit_write_task", _sync)


@pytest.fixture
def dirs(tmp_path, monkeypatch):
    """Set up temporary storage directories."""
    out = {}
    for name in ("BOMS_DIR", "QUOTES_DIR", "EPIFS_DIR"):
        d = str(tmp_path / name)
        os.makedirs(d)
        monkeypatch.setattr(config, name, d)
        out[name] = d
    template_path = os.path.join(os.path.dirname(__file__), "fixtures", "EPIF_TEMPLATE_HIRST.pdf")
    monkeypatch.setattr(config, "EPIF_TEMPLATE_PATH", template_path)
    monkeypatch.setattr(config, "ADMIN_ALERT_CHANNEL", "C_ADMIN")
    return out


def test_approve_uses_its_card_with_quote_in_thread(temp_workbook, dirs, sync_queue, monkeypatch):
    """AC 1 & 2: Drive app.handle_req_approve_action with a quote PDF in the thread.

    - Real PDF with no form fields built using pypdf.PdfWriter() + add_blank_page()
    - Row is written with the card's item, vendor, and total
    - No chat_postMessage text starts with 'Error processing'
    - epif_parser.read_fields is never called (quote carried, not read)
    - Quote is archived in QUOTES_DIR as bom.quote_filename(row, vendor, 1)
    """
    # 1. Build a real PDF with no form fields
    writer = pypdf.PdfWriter()
    writer.add_blank_page(width=612, height=792)
    buf = io.BytesIO()
    writer.write(buf)
    quote_pdf_bytes = buf.getvalue()

    # 2. Wrap epif_parser.read_fields in a spy calling the real function
    real_read_fields = epif_parser.read_fields
    read_fields_spy = MagicMock(side_effect=real_read_fields)
    monkeypatch.setattr(epif_parser, "read_fields", read_fields_spy)

    # 3. Build card request data and button value
    quote_filename = "Quote 27732 University of Wisconsin.pdf"
    file_id = "F_QUOTE_1"
    req_data = {
        "parsed": {
            "item_description": "Precision Optical Rail 500mm",
            "total_price": 349.50,
            "vendor": "Thorlabs",
            "category": "Research/Lab Supplies (3105)",
            "category_error": None,
            "project_id": "PG000025831",
            "fund": "133",
            "delivery_room": "ERB 212",
            "purpose": "Beam steering setup",
            "payment_method": "Workday",
            "date_of_purchase": "2026-09-17",
            "link": "https://thorlabs.com/item123",
            "vendor_contact_email": "orders@thorlabs.com",
            "vendor_contact_name": "Thorlabs Rep",
            "name_of_system": "",
            "asset_id": "",
        },
        "requester": "Alex",
        "user_id": "U_REQ",
        "is_pending_name": False,
        "source": "modal",
        "thread_ts": "100.0",
        "attachments": [
            {
                "id": file_id,
                "name": quote_filename,
                "role": "quote",
            }
        ],
    }

    card_blocks = [
        {
            "type": "actions",
            "elements": [
                {
                    "type": "button",
                    "action_id": "req_approve",
                    "value": json.dumps({
                        "state": "posted",
                        "requester": "Alex",
                        "thread_ts": "100.0",
                        "request": req_data,
                        "history": ["Posted by Alex"],
                    }),
                }
            ],
        }
    ]

    card_message = {
        "ts": "100.2",
        "user": "B_BOT",
        "text": "New Purchase Request from Alex",
        "metadata": {
            "event_type": "purchase_request",
            "event_payload": req_data,
        },
        "blocks": card_blocks,
    }

    bot_quote_message = {
        "ts": "100.3",
        "user": "B_BOT",
        "text": f"📎 Attached Quote: {quote_filename}",
        "files": [
            {
                "id": file_id,
                "name": quote_filename,
                "mimetype": "application/pdf",
                "url_private_download": "https://slack.com/fake/download",
            }
        ],
    }

    parent_message = {
        "ts": "100.0",
        "user": "U_REQ",
        "text": "I'd like to order an optical rail.",
    }

    # 4. Set up fake Slack client
    client = MagicMock()
    posted_messages = []

    def fake_post_message(**kwargs):
        posted_messages.append(kwargs)
        return {"ok": True, "ts": "100.99"}

    client.chat_postMessage.side_effect = fake_post_message

    def fake_replies(channel, ts, **kwargs):
        return {
            "ok": True,
            "messages": [parent_message, card_message, bot_quote_message],
        }

    client.conversations_replies.side_effect = fake_replies
    client.files_info.return_value = {
        "ok": True,
        "file": {
            "id": file_id,
            "name": quote_filename,
            "url_private_download": "https://slack.com/fake/download",
        },
    }
    client.conversations_open.return_value = {"ok": True, "channel": {"id": "D_BUYER"}}

    # Mock download_file / download to return our real blank PDF bytes
    monkeypatch.setattr(slack_io, "download_file", lambda f: quote_pdf_bytes)
    monkeypatch.setattr(slack_io, "download", lambda f: quote_pdf_bytes)

    # 5. Drive app.handle_req_approve_action
    ack = MagicMock()
    respond = MagicMock()
    body = {
        "user": {"id": "U_CHARLIE"},
        "channel": {"id": "C_PURCHASING"},
        "message": {"ts": "100.2"},
        "container": {"thread_ts": "100.0"},
        "actions": [
            {
                "action_id": "req_approve",
                "value": json.dumps({
                    "state": "posted",
                    "requester": "Alex",
                    "thread_ts": "100.0",
                    "request": req_data,
                    "history": ["Posted by Alex"],
                }),
            }
        ],
    }

    app.handle_req_approve_action(ack, body, respond, client)

    ack.assert_called_once()

    # 6. Assertions
    # epif_parser.read_fields must NEVER be called (quote is carried, not read)
    assert read_fields_spy.call_count == 0, "epif_parser.read_fields must not be called!"

    # No chat_postMessage text starts with 'Error processing'
    for msg in posted_messages:
        text = msg.get("text", "")
        assert not text.startswith("Error processing"), f"Unexpected error message posted: {text}"

    # Assert row is written with card's item, vendor, and total
    wb = openpyxl.load_workbook(temp_workbook)
    ws = wb.active
    logged_row = None
    for r in range(2, 50):
        if ws[f"C{r}"].value == "Precision Optical Rail 500mm":
            logged_row = r
            break

    assert logged_row is not None, "Expected row for 'Precision Optical Rail 500mm' was not written to workbook!"
    assert ws[f"I{logged_row}"].value == "Thorlabs"
    assert float(ws[f"{config.COLUMN_TOTAL_PRICE}{logged_row}"].value) == 349.50

    # Assert quote is archived in QUOTES_DIR as bom.quote_filename(row, vendor, 1)
    expected_quote_filename = bom.quote_filename(logged_row, "Thorlabs", 1)
    quote_path = os.path.join(dirs["QUOTES_DIR"], expected_quote_filename)
    assert os.path.exists(quote_path), f"Archived quote file {expected_quote_filename} not found in QUOTES_DIR!"
    with open(quote_path, "rb") as f:
        archived_bytes = f.read()
    assert archived_bytes == quote_pdf_bytes, "Archived quote bytes did not match expected PDF bytes!"
