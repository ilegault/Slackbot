"""Tests for Ticket 103: Quotes and a BOM dropped in the thread go with the request at approval.

Acceptance criteria:
- [ ] New tests/test_103_thread_files_at_approval.py. A thread holds a person's EPIF, a person's form-less
      Quote_A.pdf, and a person's order.xlsx (real bytes for each). Keyword approval writes the row, the
      temp QUOTES_DIR holds bom.quote_filename(row, vendor, 1), the temp BOMS_DIR holds the archived BOM
      with the .xlsx extension, and the BOM file is byte-identical to what was posted (carried, not read).
      Copy the folder setup from tests/test_83_approval_archives_attachments.py.
- [ ] Two BOM files in the thread: the exact refusal text is posted with both names, no row is written
      (the workbook row count is unchanged), and nothing is archived.
- [ ] A .png screenshot in the thread is not archived, and the attachments passed to
      finalize_purchase_request contain no entry for it.
- [ ] A quote posted on a <@BOT> quote message is not archived a second time at approval. Exactly one file
      exists in the temp QUOTES_DIR for it: the one handle_quote saved.
- [ ] A bot-posted PDF in the thread is not archived as a quote.
"""
import datetime
import io
import json
import os
import zipfile
from unittest.mock import MagicMock

import openpyxl
import pytest
from pypdf import PdfWriter

from src import (
    bom,
    config,
    epif_filler,
    lifecycle,
    log_writer,
    queue_worker,
    roster,
    slack_io,
)

FIXTURES = os.path.join(os.path.dirname(os.path.abspath(__file__)), "fixtures")
EPIF_TEMPLATE = os.path.join(FIXTURES, "EPIF_TEMPLATE_HIRST.pdf")


@pytest.fixture(autouse=True)
def clean_roster(tmp_path, monkeypatch):
    roster_file = str(tmp_path / "roster.json")
    monkeypatch.setattr(roster, "ROSTER_PATH", roster_file)
    roster._DATA = None
    data = {
        "admins": ["U_ADMIN"],
        "approvers": ["U_CHARLIE"],
        "buyers": ["U_BUYER"],
        "requesters": {
            "U_ADMIN": "Isaac",
            "U_CHARLIE": "Charlie H.",
            "U_BUYER": "Dylan",
            "U_REQ": "Alex",
        },
        "vendors": [],
    }
    with open(roster_file, "w", encoding="utf-8") as f:
        json.dump(data, f)
    return roster_file


@pytest.fixture
def dirs(tmp_path, monkeypatch):
    out = {}
    for name in ("BOMS_DIR", "QUOTES_DIR", "EPIFS_DIR"):
        d = str(tmp_path / name)
        os.makedirs(d)
        monkeypatch.setattr(config, name, d)
        out[name] = d
    monkeypatch.setattr(config, "EPIF_TEMPLATE_PATH", EPIF_TEMPLATE)
    monkeypatch.setattr(config, "ADMIN_ALERT_CHANNEL", "C_ADMIN")
    return out


@pytest.fixture
def temp_workbook(tmp_path, monkeypatch):
    from tests.test_29_approval_archives_bom import create_test_workbook
    p = str(tmp_path / "Purchasing-Log.xlsx")
    create_test_workbook(p)
    monkeypatch.setattr(config, "WORKBOOK_PATH", p)
    return p


@pytest.fixture
def sync_queue(monkeypatch):
    def _sync(action_fn, channel="", thread_ts="", user_id="", task_type="append",
              description="", success_callback=None, failure_callback=None, client=None):
        try:
            res = action_fn()
        except Exception as exc:
            if failure_callback:
                failure_callback(exc)
            return
        if success_callback:
            success_callback(res)
    monkeypatch.setattr(queue_worker, "submit_write_task", _sync)


def _row(temp_workbook):
    with zipfile.ZipFile(temp_workbook) as zf:
        xml = zf.read(config.SHEET_XML).decode("utf-8")
    return xml, log_writer.find_first_empty_row(xml) - 1


def _make_epif_bytes() -> bytes:
    with open(EPIF_TEMPLATE, "rb") as f:
        template_bytes = f.read()
    epif_data = {
        "item_description": "Tube fittings",
        "purpose": "Loop build",
        "total_price": 150.0,
        "vendor": "Swagelok",
        "vendor_contact_name": "Sales",
        "vendor_contact_email": "sales@swagelok.com",
        "date_of_purchase": datetime.date(2026, 9, 22),
        "project_id": "PG000025831",
        "fund": "133",
        "category": "Research/Lab Supplies (3105)",
        "delivery_room": "ERB 212",
        "payment_method": "P-card",
        "link": "https://swagelok.com",
        "name_of_system": "",
        "asset_id": "",
    }
    return epif_filler.fill_epif(template_bytes, epif_data)


def _make_formless_pdf_bytes() -> bytes:
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


def _make_xlsx_bytes() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"], ws["B1"] = "Part", "Price"
    ws["A2"], ws["B2"] = "SS-400-1-2", 150
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


PNG_BYTES = b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01\x08\x06\x00\x00\x00\x1f\x15c4\x00\x00\x00\nIDATx\x9cc\x00\x01\x00\x00\x05\x00\x01\r\n-\xb4\x00\x00\x00\x00IEND\xaeB`\x82"


# ---------------------------------------------------------------------------
# Criterion 1: Thread files (EPIF, Quote, BOM) go with the request at approval
# ---------------------------------------------------------------------------

def test_thread_files_archived_at_keyword_approval(temp_workbook, dirs, sync_queue, monkeypatch):
    """A thread holds a person's EPIF, form-less Quote_A.pdf, and order.xlsx (real bytes).

    Keyword approval writes the row, QUOTES_DIR holds bom.quote_filename(row, vendor, 1),
    BOMS_DIR holds the archived BOM with .xlsx, and the BOM file is byte-identical.
    """
    epif_bytes = _make_epif_bytes()
    quote_bytes = _make_formless_pdf_bytes()
    bom_bytes = _make_xlsx_bytes()

    epif_file = {"id": "F_EPIF", "name": "Swagelok_EPIF.pdf"}
    quote_file = {"id": "F_QUOTE", "name": "Quote_A.pdf"}
    bom_file = {"id": "F_BOM", "name": "order.xlsx"}

    file_bytes_map = {
        "F_EPIF": epif_bytes,
        "F_QUOTE": quote_bytes,
        "F_BOM": bom_bytes,
    }

    client = MagicMock()
    client.auth_test.return_value = {"ok": True, "user_id": "B_BOT"}
    client.conversations_open.return_value = {"channel": {"id": "D_BUYER"}}
    client.files_info.side_effect = lambda file: {"file": {"id": file, "name": epif_file["name"] if file == "F_EPIF" else (quote_file["name"] if file == "F_QUOTE" else bom_file["name"])}}

    client.conversations_replies.return_value = {
        "messages": [
            {"user": "U_REQ", "ts": "100.0", "files": [epif_file]},
            {"user": "U_REQ", "ts": "101.0", "files": [quote_file]},
            {"user": "U_REQ", "ts": "102.0", "files": [bom_file]},
        ]
    }

    monkeypatch.setattr(slack_io, "download", lambda f: file_bytes_map[f["id"]])
    monkeypatch.setattr(slack_io, "download_file", lambda f: file_bytes_map[f["id"]])
    monkeypatch.setattr(slack_io, "resolve_requester", lambda cl, uid: "Alex" if uid == "U_REQ" else "Charlie H.")
    monkeypatch.setattr(slack_io, "find_card_in_thread", lambda *a, **kw: (None, None, [], None))

    say = MagicMock()
    lifecycle.handle_epif_processing(
        client=client,
        say=say,
        channel="C_PURCHASING",
        thread_ts="100.0",
        approver="U_CHARLIE",
        event_ts="103.0",
    )

    xml, row = _row(temp_workbook)
    assert row == 17
    vendor = "Swagelok"

    # QUOTES_DIR holds bom.quote_filename(row, vendor, 1)
    expected_quote_name = bom.quote_filename(row, vendor, 1)
    assert expected_quote_name in os.listdir(dirs["QUOTES_DIR"])
    with open(os.path.join(dirs["QUOTES_DIR"], expected_quote_name), "rb") as f:
        assert f.read() == quote_bytes

    # BOMS_DIR holds archived BOM with .xlsx, byte-identical
    expected_bom_name = f"{row:04d}_{vendor}_BOM.xlsx"
    assert expected_bom_name in os.listdir(dirs["BOMS_DIR"])
    with open(os.path.join(dirs["BOMS_DIR"], expected_bom_name), "rb") as f:
        assert f.read() == bom_bytes

    # Notes cell has BOM record
    assert log_writer.get_cell_value(xml, f"{config.COLUMN_NOTES}{row}") == f"BOM: {expected_bom_name} (attached)"


# ---------------------------------------------------------------------------
# Criterion 2: Two BOM files in thread refused with exact wording
# ---------------------------------------------------------------------------

def test_two_boms_in_thread_refused_no_row_written(temp_workbook, dirs, sync_queue, monkeypatch):
    """Two BOM files in the thread: exact refusal text is posted with both names,

    no row is written (workbook row count unchanged), and nothing is archived.
    """
    epif_bytes = _make_epif_bytes()
    bom1_bytes = _make_xlsx_bytes()
    bom2_bytes = b"part,qty\n123,1\n"

    epif_file = {"id": "F_EPIF", "name": "Swagelok_EPIF.pdf"}
    bom1_file = {"id": "F_BOM1", "name": "order1.xlsx"}
    bom2_file = {"id": "F_BOM2", "name": "order2.csv"}

    client = MagicMock()
    client.auth_test.return_value = {"ok": True, "user_id": "B_BOT"}
    client.conversations_replies.return_value = {
        "messages": [
            {"user": "U_REQ", "ts": "100.0", "files": [epif_file]},
            {"user": "U_REQ", "ts": "101.0", "files": [bom1_file]},
            {"user": "U_REQ", "ts": "102.0", "files": [bom2_file]},
        ]
    }

    file_bytes_map = {"F_EPIF": epif_bytes, "F_BOM1": bom1_bytes, "F_BOM2": bom2_bytes}
    monkeypatch.setattr(slack_io, "download", lambda f: file_bytes_map[f["id"]])
    monkeypatch.setattr(slack_io, "download_file", lambda f: file_bytes_map[f["id"]])
    monkeypatch.setattr(slack_io, "resolve_requester", lambda cl, uid: "Alex")
    monkeypatch.setattr(slack_io, "find_card_in_thread", lambda *a, **kw: (None, None, [], None))

    say = MagicMock()
    lifecycle.handle_epif_processing(
        client=client,
        say=say,
        channel="C_PURCHASING",
        thread_ts="100.0",
        approver="U_CHARLIE",
        event_ts="103.0",
    )

    # Exact refusal text
    say.assert_called_once()
    refusal_text = say.call_args[1]["text"]
    assert refusal_text == "One BOM per EPIF \u2014 delete the extra and approve again: `order1.xlsx`, `order2.csv`"

    # No row written: row count remains 16 (empty row is still 17)
    xml, row = _row(temp_workbook)
    assert row == 16

    # Nothing archived
    assert os.listdir(dirs["BOMS_DIR"]) == []
    assert os.listdir(dirs["QUOTES_DIR"]) == []
    assert os.listdir(dirs["EPIFS_DIR"]) == []


# ---------------------------------------------------------------------------
# Criterion 3: Screenshot .png in thread ignored
# ---------------------------------------------------------------------------

def test_screenshot_png_in_thread_not_archived(temp_workbook, dirs, sync_queue, monkeypatch):
    """A .png screenshot in the thread is not archived, and attachments passed to

    finalize_purchase_request contain no entry for it.
    """
    epif_bytes = _make_epif_bytes()
    quote_bytes = _make_formless_pdf_bytes()

    epif_file = {"id": "F_EPIF", "name": "Swagelok_EPIF.pdf"}
    png_file = {"id": "F_PNG", "name": "screenshot.png"}
    quote_file = {"id": "F_QUOTE", "name": "Quote_A.pdf"}

    file_bytes_map = {
        "F_EPIF": epif_bytes,
        "F_PNG": PNG_BYTES,
        "F_QUOTE": quote_bytes,
    }

    client = MagicMock()
    client.auth_test.return_value = {"ok": True, "user_id": "B_BOT"}
    client.conversations_open.return_value = {"channel": {"id": "D_BUYER"}}
    client.files_info.side_effect = lambda file: {"file": {"id": file, "name": quote_file["name"] if file == "F_QUOTE" else "file"}}

    client.conversations_replies.return_value = {
        "messages": [
            {"user": "U_REQ", "ts": "100.0", "files": [epif_file]},
            {"user": "U_REQ", "ts": "101.0", "files": [png_file]},
            {"user": "U_REQ", "ts": "102.0", "files": [quote_file]},
        ]
    }

    monkeypatch.setattr(slack_io, "download", lambda f: file_bytes_map[f["id"]])
    monkeypatch.setattr(slack_io, "download_file", lambda f: file_bytes_map[f["id"]])
    monkeypatch.setattr(slack_io, "resolve_requester", lambda cl, uid: "Alex")
    monkeypatch.setattr(slack_io, "find_card_in_thread", lambda *a, **kw: (None, None, [], None))

    captured_attachments = []
    orig_finalize = lifecycle.finalize_purchase_request

    def spy_finalize(*args, **kwargs):
        captured_attachments.append(kwargs.get("attachments"))
        return orig_finalize(*args, **kwargs)

    monkeypatch.setattr(lifecycle, "finalize_purchase_request", spy_finalize)

    say = MagicMock()
    lifecycle.handle_epif_processing(
        client=client,
        say=say,
        channel="C_PURCHASING",
        thread_ts="100.0",
        approver="U_CHARLIE",
        event_ts="103.0",
    )

    assert len(captured_attachments) == 1
    atts = captured_attachments[0]
    assert atts is not None
    # No screenshot entry in attachments
    assert not any(a.get("name") == "screenshot.png" or a.get("id") == "F_PNG" for a in atts)
    # Only quote in attachments
    assert atts == [{"id": "F_QUOTE", "name": "Quote_A.pdf", "role": "quote"}]

    # Only quote archived in QUOTES_DIR, no screenshot
    assert os.listdir(dirs["QUOTES_DIR"]) == ["0017_Swagelok_Quote_1.pdf"]


# ---------------------------------------------------------------------------
# Criterion 4: Quote on @BOT quote message not archived a second time
# ---------------------------------------------------------------------------

def test_quote_on_bot_quote_message_not_archived_again(temp_workbook, dirs, sync_queue, monkeypatch):
    """A quote posted on a <@BOT> quote message is not archived a second time at approval.

    Exactly one file exists in the temp QUOTES_DIR for it: the one handle_quote saved.
    """
    epif_bytes = _make_epif_bytes()
    quote_bytes = _make_formless_pdf_bytes()

    epif_file = {"id": "F_EPIF", "name": "Swagelok_EPIF.pdf"}
    quote_file = {"id": "F_QUOTE", "name": "Vendor_Quote.pdf"}

    file_bytes_map = {
        "F_EPIF": epif_bytes,
        "F_QUOTE": quote_bytes,
    }

    client = MagicMock()
    client.auth_test.return_value = {"ok": True, "user_id": "B_BOT"}
    client.conversations_open.return_value = {"channel": {"id": "D_BUYER"}}
    monkeypatch.setattr(slack_io, "download", lambda f: file_bytes_map[f["id"]])
    monkeypatch.setattr(slack_io, "download_file", lambda f: file_bytes_map[f["id"]])
    monkeypatch.setattr(slack_io, "resolve_requester", lambda cl, uid: "Alex")
    monkeypatch.setattr(slack_io, "find_card_in_thread", lambda *a, **kw: (None, None, [], None))

    # 1. Requester uploads quote with @Purchasing quote command
    say_quote = MagicMock()
    lifecycle.handle_quote(
        client=client,
        say=say_quote,
        channel="C_PURCHASING",
        thread_ts="100.0",
        event_ts="101.0",
        files=[quote_file],
    )
    # Exactly one file saved by handle_quote
    assert os.listdir(dirs["QUOTES_DIR"]) == ["Vendor_Quote.pdf"]

    # 2. Later, approver approves the thread. Thread replies contain the quote command message.
    client.conversations_replies.return_value = {
        "messages": [
            {"user": "U_REQ", "ts": "100.0", "files": [epif_file]},
            {"user": "U_REQ", "ts": "101.0", "text": "<@B_BOT> quote", "files": [quote_file]},
        ]
    }

    say_appr = MagicMock()
    lifecycle.handle_epif_processing(
        client=client,
        say=say_appr,
        channel="C_PURCHASING",
        thread_ts="100.0",
        approver="U_CHARLIE",
        event_ts="102.0",
    )

    # Row was logged
    _, row = _row(temp_workbook)
    assert row == 17

    # Still exactly one file in QUOTES_DIR: the one handle_quote saved!
    assert os.listdir(dirs["QUOTES_DIR"]) == ["Vendor_Quote.pdf"]


# ---------------------------------------------------------------------------
# Criterion 5: Bot-posted PDF in thread not archived as quote
# ---------------------------------------------------------------------------

def test_bot_posted_pdf_not_archived_as_quote(temp_workbook, dirs, sync_queue, monkeypatch):
    """A bot-posted PDF in the thread is not archived as a quote."""
    epif_bytes = _make_epif_bytes()
    bot_pdf_bytes = _make_formless_pdf_bytes()

    epif_file = {"id": "F_EPIF", "name": "Swagelok_EPIF.pdf"}
    bot_file = {"id": "F_BOT_PDF", "name": "bot_summary.pdf"}

    file_bytes_map = {
        "F_EPIF": epif_bytes,
        "F_BOT_PDF": bot_pdf_bytes,
    }

    client = MagicMock()
    client.auth_test.return_value = {"ok": True, "user_id": "B_BOT"}
    client.conversations_open.return_value = {"channel": {"id": "D_BUYER"}}

    client.conversations_replies.return_value = {
        "messages": [
            {"user": "U_REQ", "ts": "100.0", "files": [epif_file]},
            {"user": "B_BOT", "ts": "101.0", "files": [bot_file]},
        ]
    }

    monkeypatch.setattr(slack_io, "download", lambda f: file_bytes_map[f["id"]])
    monkeypatch.setattr(slack_io, "download_file", lambda f: file_bytes_map[f["id"]])
    monkeypatch.setattr(slack_io, "resolve_requester", lambda cl, uid: "Alex")
    monkeypatch.setattr(slack_io, "find_card_in_thread", lambda *a, **kw: (None, None, [], None))

    say = MagicMock()
    lifecycle.handle_epif_processing(
        client=client,
        say=say,
        channel="C_PURCHASING",
        thread_ts="100.0",
        approver="U_CHARLIE",
        event_ts="102.0",
    )

    _, row = _row(temp_workbook)
    assert row == 17

    # QUOTES_DIR has nothing archived
    assert os.listdir(dirs["QUOTES_DIR"]) == []


# ---------------------------------------------------------------------------
# Unit test: slack_io.thread_files and find_epif_in_thread delegation
# ---------------------------------------------------------------------------

def test_thread_files_and_find_epif_in_thread_delegation(monkeypatch):
    """slack_io.thread_files classifies files; find_epif_in_thread delegates to it."""
    client = MagicMock()
    client.auth_test.return_value = {"ok": True, "user_id": "B_BOT"}

    epif_file = {"id": "F1", "name": "form.pdf"}
    quote_file = {"id": "F2", "name": "quote.pdf"}

    client.conversations_replies.return_value = {
        "messages": [
            {"user": "U_REQ", "ts": "100.0", "files": [epif_file]},
            {"user": "U_REQ", "ts": "101.0", "files": [quote_file]},
        ]
    }

    def fake_download(file_obj):
        if file_obj.get("id") == "F1":
            return _make_epif_bytes()
        return _make_formless_pdf_bytes()

    monkeypatch.setattr(slack_io, "download", fake_download)

    res = slack_io.thread_files(client, "C123", "100.0")
    assert len(res["epifs"]) == 1
    assert res["epifs"][0]["file"] == epif_file
    assert len(res["quotes"]) == 1
    assert res["quotes"][0]["file"] == quote_file

    # find_epif_in_thread delegates to thread_files and returns newest epif
    found_file, found_user = slack_io.find_epif_in_thread(client, "C123", "100.0")
    assert found_file == epif_file
    assert found_user == "U_REQ"
