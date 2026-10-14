"""Ticket 132: the rebuilt BOM follows the lab template, reads back, and is named by the lab."""
import datetime
import io
import os

import openpyxl

from src import bom, log_writer

FIXTURE = os.path.join(os.path.dirname(__file__), "fixtures", "BOM_TEMPLATE_FILLED_EXAMPLE.xlsx")


def test_rebuilt_bom_reads_back():
    with open(FIXTURE, "rb") as f:
        items, problems = bom.read_bom_template(f.read())
    assert problems == [] and len(items) == 4

    request = {"item_description": "Duniway Stockroom Corp. Copper Gaskets", "requester": "A Lab Member"}
    rebuilt = bom.build_bom_workbook(request, items)
    again, problems2 = bom.read_bom_template(rebuilt)
    assert problems2 == []
    assert again == items  # includes links
    assert all(i["link"].startswith("https://www.duniway.com/") for i in again)

    ws = openpyxl.load_workbook(io.BytesIO(rebuilt)).worksheets[0]
    assert ws["B2"].value == "Duniway Stockroom Corp. Copper Gaskets"
    assert ws["B3"].value == "A Lab Member"
    template_ws = openpyxl.load_workbook(bom.config.BOM_TEMPLATE_PATH).worksheets[0]
    header_row = next(r for r in range(1, 11) if template_ws.cell(r, 1).value == "Item")
    assert ws.cell(header_row, 1).value == "Item" and ws.cell(header_row, 8).value == "Notes"
    assert bom.check_bom_total(again, 858.50) is None


def test_bom_filename_cases():
    d = datetime.date(2026, 10, 14)
    assert (bom.bom_filename("Copper gaskets for CF flanges", d)
            == "HIRST_LAB_BOM_Copper-gaskets-for-CF-flanges_10.14.2026.xlsx")
    name = bom.bom_filename('a/b:c*d?e"f<g>h|i', d)
    assert name == "HIRST_LAB_BOM_abcdefghi_10.14.2026.xlsx"
    long_name = bom.bom_filename("x" * 200, d)
    assert long_name == "HIRST_LAB_BOM_" + "x" * 60 + "_10.14.2026.xlsx"
    assert bom.bom_filename("Gaskets", d, draft=True).startswith("DRAFT_HIRST_LAB_BOM_Gaskets_")


def test_save_bom_never_overwrites(tmp_path):
    name = "HIRST_LAB_BOM_Gaskets_10.14.2026.xlsx"
    first = log_writer.save_bom(b"one", name, target_dir=str(tmp_path))
    second = log_writer.save_bom(b"two", name, target_dir=str(tmp_path))
    assert os.path.basename(first) == name
    assert os.path.basename(second) == "HIRST_LAB_BOM_Gaskets_10.14.2026_2.xlsx"
    assert sorted(os.listdir(tmp_path)) == sorted([name, os.path.basename(second)])
    assert open(first, "rb").read() == b"one" and open(second, "rb").read() == b"two"
