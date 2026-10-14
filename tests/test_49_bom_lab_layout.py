"""Tickets 49 and 132: the BOM sheet follows the lab template (ADR 0019). Real bytes, nothing faked."""
import io
import re

import openpyxl

from src import bom

ITEMS = [
    {"qty": 2, "name": "Shaft Collar", "part_number": "SC-100", "unit_price": 12.50,
     "link": "https://ruland.com/sc100", "description": "Steel collar"},
    {"qty": 4, "name": "Clamp Collar", "part_number": "CC-200", "unit_price": 15.00,
     "link": "", "description": "Clamp"},
    {"qty": 3, "name": "Coupling", "part_number": "", "unit_price": 7.25,
     "link": "see quote", "description": ""},
]
HEADERS = ["Item", "Description ", "Product #", "Vendor", "Unit Cost", "Quantity", "Total Cost", "Notes"]


def _open(request, items):
    return openpyxl.load_workbook(io.BytesIO(bom.build_bom_workbook(request, items))).active


def _eval(ws, coord):
    """Evaluate only the three formula shapes this sheet uses; raise on anything else."""
    v = ws[coord].value
    if not (isinstance(v, str) and v.startswith("=")):
        return v
    m = re.fullmatch(r"=F(\d+)\*E(\d+)", v)
    if m and m.group(1) == m.group(2):
        return _eval(ws, f"E{m.group(1)}") * _eval(ws, f"F{m.group(1)}")
    m = re.fullmatch(r"=SUM\(G(\d+):G(\d+)\)", v)
    if m:
        return sum(_eval(ws, f"G{r}") for r in range(int(m.group(1)), int(m.group(2)) + 1))
    m = re.fullmatch(r"=G(\d+)", v)
    if m:
        return _eval(ws, f"G{m.group(1)}")
    raise AssertionError(f"unexpected formula {v!r} in {coord}")


def test_layout():
    ws = _open({"vendor": "Ruland", "requester": "A Buyer", "item_description": "Collars"}, ITEMS)
    assert ws.title == "BOM"
    assert ws["A1"].value == "BILL OF MATERIALS"
    assert "A1:B1" in [str(r) for r in ws.merged_cells.ranges]
    assert (ws["A2"].value, ws["B2"].value) == ("Order Description", "Collars")
    assert (ws["A3"].value, ws["B3"].value) == ("BOM Creator", "A Buyer")
    assert ws["A4"].value == "Total Cost"
    assert [c.value for c in ws[5]][:8] == HEADERS
    for i, it in enumerate(ITEMS):
        r = 6 + i
        # openpyxl reads an empty-string cell back as None
        assert [ws.cell(r, c).value for c in range(1, 7)] == [
            it["name"], it["description"] or None, it["part_number"] or None, "Ruland",
            it["unit_price"], it["qty"]]
        assert ws.cell(r, 8).value is None
    assert ws["F9"].value == "Total"
    assert ws.freeze_panes == "A6"
    assert ws["A6"].hyperlink is not None and ws["A6"].hyperlink.target == ITEMS[0]["link"]
    assert ws["A7"].hyperlink is None and ws["A8"].hyperlink is None
    # The template's own styling survives (not a look-alike built from scratch).
    assert ws["G6"].number_format == ws["G8"].number_format
    assert ws["F9"].font.b


def test_formulas_compute_right_total():
    ws = _open({"vendor": "Ruland"}, ITEMS)
    for i, it in enumerate(ITEMS):
        assert _eval(ws, f"G{6 + i}") == it["qty"] * it["unit_price"]
    expected = sum(it["qty"] * it["unit_price"] for it in ITEMS)
    assert _eval(ws, "G9") == expected
    assert _eval(ws, "B4") == expected


def test_unknown_vendor_fallback():
    ws = _open({"vendor": "  "}, ITEMS)
    assert ws["D6"].value == "Unknown Vendor"


def test_nothing_the_adr_removed():
    request = {"vendor": "Ruland", "requester": "ZZrequester", "project_id": "PROJ-9917",
               "fund": "FUND-4471", "date_of_purchase": "2031-02-03"}
    items = [{"qty": i + 1, "name": f"Part {i}", "part_number": f"P{i}", "unit_price": 1.0 + i,
              "link": "", "description": ""} for i in range(7)]
    ws = _open(request, items)
    for row in ws.iter_rows():
        for c in row:
            if c.value is None:
                continue
            text = str(c.value)
            # The requester name now belongs in B3 (ADR 0019); nowhere else.
            if c.coordinate != "B3":
                assert "ZZrequester" not in text, c.coordinate
            for banned in ("PROJ-9917", "FUND-4471", "2031-02-03",
                           "Shipping", "DRAFT", "Purchasing Log"):
                assert banned not in text, f"{c.coordinate} contains {banned}"
    assert ws["B3"].value == "ZZrequester"
    # More items than the template's four rows: the sheet grows, total row follows the items.
    last = max(c.row for row in ws.iter_rows() for c in row if c.value is not None)
    assert last == 13 and ws["F13"].value == "Total"
