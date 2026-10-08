"""Tests for Ticket 105: A batch refuses a BOM and carries the thread's quotes on every row.

Acceptance criteria:
- [ ] New tests/test_105_batch_files.py: a two-EPIF thread plus a person's order.xlsx.
      The exact refusal text is posted with n = 2, no row is written, and the temp BOMS_DIR
      and EPIFS_DIR are empty.
- [ ] A two-EPIF thread plus two person-posted quote PDFs. The temp QUOTES_DIR holds four files,
      quote_filename(row_A, vendor_A, 1..2) and quote_filename(row_B, vendor_B, 1..2),
      each byte-identical to the posted quote.
- [ ] Assigning a buyer to card B (drive lifecycle.handle_assign as tests/test_30_buyers_dm_carries_bom.py does)
      uploads card B's two quote paths in the buyer's DM. Assert on the files_upload_v2 calls' file paths.
- [ ] A single-EPIF thread with a BOM is unaffected. Ticket 103's tests pass unchanged.

May fake: the Slack client, downloads (return fixture bytes), parse_epif.
Must be real: the batch path, finalize, archive writes into temp folders, the workbook write on a temp copy.
"""
import datetime
import io
import json
import os
import zipfile
from unittest.mock import MagicMock, patch

import openpyxl
import pytest
from pypdf import PdfWriter

from src import (
    bom,
    config,
    epif_parser,
    lifecycle,
    log_writer,
    queue_worker,
    roster,
    slack_io,
)

FIXTURES = os.path.join(os.path.dirname(os.path.abspath(__file__)), "fixtures")
EPIF_TEMPLATE = os.path.join(FIXTURES, "EPIF_TEMPLATE_HIRST.pdf")


@pytest.fixture(autouse=True)
def clean_roster(tmp_path, monkeypatch):
    roster_file = str(tmp_path / "roster.json")
    monkeypatch.setattr(roster, "ROSTER_PATH", roster_file)
    roster._DATA = None
    data = {
        "admins": ["U_ADMIN"],
        "approvers": ["U_CHARLIE"],
        "buyers": ["U_BUYER"],
        "requesters": {
            "U_ADMIN": "Isaac",
            "U_CHARLIE": "Charlie H.",
            "U_BUYER": "Dylan",
            "U_REQ": "Alex",
        },
        "vendors": [],
    }
    with open(roster_file, "w", encoding="utf-8") as f:
        json.dump(data, f)
    return roster_file


@pytest.fixture
def dirs(tmp_path, monkeypatch):
    out = {}
    for name in ("BOMS_DIR", "QUOTES_DIR", "EPIFS_DIR"):
        d = str(tmp_path / name)
        os.makedirs(d)
        monkeypatch.setattr(config, name, d)
        out[name] = d
    monkeypatch.setattr(config, "EPIF_TEMPLATE_PATH", EPIF_TEMPLATE)
    monkeypatch.setattr(config, "ADMIN_ALERT_CHANNEL", "C_ADMIN")
    return out


@pytest.fixture
def temp_workbook(tmp_path, monkeypatch):
    from tests.test_29_approval_archives_bom import create_test_workbook

    p = str(tmp_path / "Purchasing-Log.xlsx")
    create_test_workbook(p)
    monkeypatch.setattr(config, "WORKBOOK_PATH", p)
    return p


@pytest.fixture
def sync_queue(monkeypatch):
    def _sync(
        action_fn,
        channel="",
        thread_ts="",
        user_id="",
        task_type="append",
        description="",
        success_callback=None,
        failure_callback=None,
        client=None,
    ):
        try:
            res = action_fn()
        except Exception as exc:
            if failure_callback:
                failure_callback(exc)
            return
        if success_callback:
            success_callback(res)

    monkeypatch.setattr(queue_worker, "submit_write_task", _sync)


def _row(temp_workbook):
    with zipfile.ZipFile(temp_workbook) as zf:
        xml = zf.read(config.SHEET_XML).decode("utf-8")
    return xml, log_writer.find_first_empty_row(xml) - 1


def _make_parsed(vendor: str, total_price: float, item_desc: str = "Test Supplies"):
    return {
        "item_description": item_desc,
        "purpose": "Experiment setup",
        "total_price": total_price,
        "vendor": vendor,
        "vendor_contact_name": "Sales",
        "vendor_contact_email": f"sales@{vendor.lower().replace(' ', '')}.com",
        "date_of_purchase": datetime.date(2026, 9, 22),
        "project_id": "PG000025831",
        "fund": "133",
        "category": "Research/Lab Supplies (3105)",
        "delivery_room": "ERB 212",
        "payment_method": "P-card",
        "link": f"https://{vendor.lower().replace(' ', '')}.com",
        "name_of_system": "",
        "asset_id": "",
    }


def _make_xlsx_bytes() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"], ws["B1"] = "Part", "Price"
    ws["A2"], ws["B2"] = "SS-400-1-2", 150
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _make_formless_pdf_bytes() -> bytes:
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Criterion 1: Two-EPIF thread plus order.xlsx refuses approval
# ---------------------------------------------------------------------------

def test_batch_with_bom_refuses_approval(temp_workbook, dirs, sync_queue, monkeypatch):
    """A two-EPIF thread plus a person's order.xlsx posts the exact refusal text with n = 2,
    writes no rows, and leaves temp BOMS_DIR and EPIFS_DIR empty.
    """
    _, initial_row = _row(temp_workbook)

    file_a = {"id": "F_A", "name": "VendorA_EPIF.pdf"}
    file_b = {"id": "F_B", "name": "VendorB_EPIF.pdf"}
    file_bom = {"id": "F_BOM", "name": "order.xlsx"}

    bom_bytes = _make_xlsx_bytes()
    file_bytes_map = {
        "F_A": b"%PDF-A",
        "F_B": b"%PDF-B",
        "F_BOM": bom_bytes,
    }

    client = MagicMock()
    client.auth_test.return_value = {"ok": True, "user_id": "B_BOT"}
    client.conversations_replies.return_value = {
        "messages": [
            {"user": "U_REQ", "ts": "100.0", "files": [file_a]},
            {"user": "U_REQ", "ts": "101.0", "files": [file_b]},
            {"user": "U_REQ", "ts": "102.0", "files": [file_bom]},
        ]
    }

    monkeypatch.setattr(slack_io, "download", lambda f: file_bytes_map[f["id"]])
    monkeypatch.setattr(slack_io, "download_file", lambda f: file_bytes_map[f["id"]])
    monkeypatch.setattr(slack_io, "resolve_requester", lambda cl, uid: "Alex" if uid == "U_REQ" else "Charlie H.")
    monkeypatch.setattr(slack_io, "find_card_in_thread", lambda *a, **kw: (None, None, [], None))
    monkeypatch.setattr(epif_parser, "read_fields", lambda b: {"Amount of Purchase": "1", "Vendor": "V"})
    monkeypatch.setattr(
        epif_parser,
        "parse_epif",
        lambda b: _make_parsed("Vendor A", 50.0) if b == b"%PDF-A" else _make_parsed("Vendor B", 75.0),
    )

    say = MagicMock()
    lifecycle.handle_epif_processing(
        client=client,
        say=say,
        channel="C_PURCHASING",
        thread_ts="100.0",
        event_ts="100.1",
        approver="U_CHARLIE",
    )

    # 1. Exact refusal text posted with n = 2
    say.assert_called_once()
    called_text = say.call_args[1].get("text") or say.call_args[0][0]
    expected_refusal = (
        "This thread has 2 EPIFs and a BOM — I can't tell which EPIF the BOM belongs to. "
        "Put each vendor's order in its own thread."
    )
    assert called_text == expected_refusal

    # 2. No row written
    _, after_row = _row(temp_workbook)
    assert after_row == initial_row

    # 3. Temp BOMS_DIR and EPIFS_DIR are empty
    assert os.listdir(dirs["BOMS_DIR"]) == []
    assert os.listdir(dirs["EPIFS_DIR"]) == []


# ---------------------------------------------------------------------------
# Criterion 2: Two-EPIF thread plus two quote PDFs archives 4 quote files
# ---------------------------------------------------------------------------

def test_batch_quotes_archived_per_row(temp_workbook, dirs, sync_queue, monkeypatch):
    """A two-EPIF thread plus two person-posted quote PDFs archives four files in QUOTES_DIR:
    quote_filename(row_A, vendor_A, 1..2) and quote_filename(row_B, vendor_B, 1..2),
    each byte-identical to the posted quote.
    """
    _, initial_row = _row(temp_workbook)

    file_a = {"id": "F_A", "name": "VendorA_EPIF.pdf"}
    file_b = {"id": "F_B", "name": "VendorB_EPIF.pdf"}
    file_q1 = {"id": "F_Q1", "name": "Quote1.pdf"}
    file_q2 = {"id": "F_Q2", "name": "Quote2.pdf"}

    q1_bytes = b"%PDF-QUOTE-1-BYTES"
    q2_bytes = b"%PDF-QUOTE-2-BYTES"
    file_bytes_map = {
        "F_A": b"%PDF-A",
        "F_B": b"%PDF-B",
        "F_Q1": q1_bytes,
        "F_Q2": q2_bytes,
    }

    client = MagicMock()
    client.auth_test.return_value = {"ok": True, "user_id": "B_BOT"}
    client.conversations_replies.return_value = {
        "messages": [
            {"user": "U_REQ", "ts": "100.0", "files": [file_a]},
            {"user": "U_REQ", "ts": "101.0", "files": [file_b]},
            {"user": "U_REQ", "ts": "102.0", "files": [file_q1]},
            {"user": "U_REQ", "ts": "103.0", "files": [file_q2]},
        ]
    }
    client.files_info.side_effect = lambda file: {
        "file": {"id": file, "name": file_q1["name"] if file == "F_Q1" else file_q2["name"]}
    }

    monkeypatch.setattr(slack_io, "download", lambda f: file_bytes_map[f["id"]])
    monkeypatch.setattr(slack_io, "download_file", lambda f: file_bytes_map[f["id"]])
    monkeypatch.setattr(slack_io, "resolve_requester", lambda cl, uid: "Alex" if uid == "U_REQ" else "Charlie H.")
    monkeypatch.setattr(slack_io, "find_card_in_thread", lambda *a, **kw: (None, None, [], None))
    # read_fields distinguishes EPIFs (returns non-empty fields) from quotes (returns empty dict)
    monkeypatch.setattr(
        epif_parser,
        "read_fields",
        lambda b: {"Amount of Purchase": "1", "Vendor": "V"} if b in (b"%PDF-A", b"%PDF-B") else {},
    )
    monkeypatch.setattr(
        epif_parser,
        "parse_epif",
        lambda b: _make_parsed("Vendor A", 50.0) if b == b"%PDF-A" else _make_parsed("Vendor B", 75.0),
    )

    say = MagicMock()
    lifecycle.handle_epif_processing(
        client=client,
        say=say,
        channel="C_PURCHASING",
        thread_ts="100.0",
        event_ts="100.1",
        approver="U_CHARLIE",
    )

    # Two rows written: initial_row + 1 and initial_row + 2
    row_a = initial_row + 1
    row_b = initial_row + 2
    _, after_row = _row(temp_workbook)
    assert after_row == row_b

    # Four files in QUOTES_DIR
    q_dir = dirs["QUOTES_DIR"]
    files_in_quotes = sorted(os.listdir(q_dir))
    expected_q_a1 = bom.quote_filename(row_a, "Vendor A", 1)
    expected_q_a2 = bom.quote_filename(row_a, "Vendor A", 2)
    expected_q_b1 = bom.quote_filename(row_b, "Vendor B", 1)
    expected_q_b2 = bom.quote_filename(row_b, "Vendor B", 2)
    expected_files = sorted([expected_q_a1, expected_q_a2, expected_q_b1, expected_q_b2])
    assert files_in_quotes == expected_files

    # Assert byte identity
    with open(os.path.join(q_dir, expected_q_a1), "rb") as f:
        assert f.read() == q1_bytes
    with open(os.path.join(q_dir, expected_q_a2), "rb") as f:
        assert f.read() == q2_bytes
    with open(os.path.join(q_dir, expected_q_b1), "rb") as f:
        assert f.read() == q1_bytes
    with open(os.path.join(q_dir, expected_q_b2), "rb") as f:
        assert f.read() == q2_bytes


# ---------------------------------------------------------------------------
# Criterion 3: Assigning a buyer to card B uploads card B's two quote paths
# ---------------------------------------------------------------------------

def test_assign_card_b_uploads_quotes_in_buyer_dm(temp_workbook, dirs, sync_queue, monkeypatch):
    """Assigning a buyer to card B (driven as test_30 does) uploads card B's two quote paths in the buyer's DM.
    Assert on files_upload_v2 calls' file paths.
    """
    row_b = 18
    vendor_b = "Vendor B"
    q_dir = dirs["QUOTES_DIR"]

    q1_fname = bom.quote_filename(row_b, vendor_b, 1)
    q2_fname = bom.quote_filename(row_b, vendor_b, 2)
    q1_path = os.path.join(q_dir, q1_fname)
    q2_path = os.path.join(q_dir, q2_fname)

    with open(q1_path, "wb") as f:
        f.write(b"%PDF-QUOTE-1-BYTES")
    with open(q2_path, "wb") as f:
        f.write(b"%PDF-QUOTE-2-BYTES")

    client = MagicMock()
    client.conversations_open.return_value = {"channel": {"id": "D_BUYER_DM"}}
    client.files_upload_v2 = MagicMock()

    card_b_req = {
        "item_description": "Vacuum Flanges",
        "route": "epif",
        "source": "epif",
        "vendor": vendor_b,
        "total_price": 75.0,
        "quote_count": 2,
    }
    history = ["Approved by Charlie on 09/22"]

    say = MagicMock()

    with patch.object(lifecycle.slack_io, "find_card_in_thread",
                      return_value=(card_b_req, "100.2", history, "approved")):
        with patch.object(lifecycle.slack_io, "find_row_in_thread", return_value=row_b):
            with patch.object(lifecycle.log_writer, "get_row_info", return_value={}):
                ok = lifecycle.handle_assign(
                    client=client,
                    say=say,
                    channel="C_PURCHASING",
                    thread_ts="100.0",
                    user_id="U_CHARLIE",
                    event_ts="100.3",
                    target_user_id="U_BUYER",
                )

    assert ok is True

    # Assert files_upload_v2 was called for card B's two quote paths
    client.files_upload_v2.assert_called()
    calls = client.files_upload_v2.call_args_list

    # Assert on file paths
    uploaded_files = [c[1].get("file") for c in calls if "file" in c[1]]
    assert q1_path in uploaded_files
    assert q2_path in uploaded_files

    # Assert on filenames
    uploaded_filenames = [c[1].get("filename") for c in calls]
    assert q1_fname in uploaded_filenames
    assert q2_fname in uploaded_filenames


# ---------------------------------------------------------------------------
# Criterion 4: Single-EPIF thread with a BOM is unaffected
# ---------------------------------------------------------------------------

def test_single_epif_with_bom_unaffected(temp_workbook, dirs, sync_queue, monkeypatch):
    """A single-EPIF thread with a BOM is unaffected and archives both EPIF and BOM."""
    from tests.test_103_thread_files_at_approval import test_thread_files_archived_at_keyword_approval

    test_thread_files_archived_at_keyword_approval(temp_workbook, dirs, sync_queue, monkeypatch)
