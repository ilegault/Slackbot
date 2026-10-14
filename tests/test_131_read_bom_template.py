"""Ticket 131: the pure reader for the lab BOM template (ADR 0019)."""
import csv
import io
import os

import openpyxl

from src import bom, config

FIXTURES = os.path.join(os.path.dirname(__file__), "fixtures")
FILLED = os.path.join(FIXTURES, "BOM_TEMPLATE_FILLED_EXAMPLE.xlsx")
EPIF_PDF = os.path.join(FIXTURES, "EPIF_TEMPLATE_HIRST.pdf")

NOT_TEMPLATE = (
    "This isn't the lab BOM template — get it with the “Send me the BOM template” "
    "button. Vendor carts and quotes go in the Quotes field."
)


def _filled_wb():
    return openpyxl.load_workbook(FILLED)


def _to_bytes(wb) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_reads_the_real_lab_bom():
    with open(FILLED, "rb") as fh:
        items, problems = bom.read_bom_template(fh.read())
    assert problems == []
    assert len(items) == 4
    first = items[0]
    assert first["quantity"] == 2
    assert first["unit_price"] == 23.75
    assert first["part_number"] == "G-0133"
    assert first["link"] == "https://www.duniway.com/part/g-0133"
    assert bom.bom_total(items) == 858.50


def test_not_the_template():
    wb = openpyxl.Workbook()
    wb.active.append(["Qty", "Part", "Price"])
    wb.active.append([1, "x", 2])
    wrong_headers = _to_bytes(wb)

    out = io.StringIO()
    csv.writer(out).writerow(config.BOM_TEMPLATE_HEADERS)
    csv_bytes = out.getvalue().encode()

    with open(EPIF_PDF, "rb") as fh:
        pdf_bytes = fh.read()

    for data in (wrong_headers, csv_bytes, pdf_bytes):
        assert bom.read_bom_template(data) == ([], [NOT_TEMPLATE])


def test_row_problems_by_row_number():
    wb = _filled_wb()
    ws = wb.active
    ws["E7"] = "TBD"
    ws["F8"] = 2.5
    ws["A9"] = None
    # openpyxl refills a blank cell from its hyperlink on load, so clear that too.
    ws["A9"].hyperlink = None
    ws.insert_rows(10)  # a fully blank row above the stray mark / Total row
    items, problems = bom.read_bom_template(_to_bytes(wb))
    assert "Row 7: Unit Cost 'TBD' isn't a number." in problems
    assert "Row 8: Quantity '2.5' isn't a whole number of 1 or more." in problems
    assert "Row 9: Item is blank." in problems
    assert len(problems) == 3  # blank row and the stray backtick make no problem
    assert len(items) == 1  # only row 6 read cleanly


def test_check_bom_total():
    with open(FILLED, "rb") as fh:
        items, _ = bom.read_bom_template(fh.read())
    assert bom.check_bom_total(items, 858.50) is None
    assert bom.check_bom_total(items, 858.49) == (
        "The BOM totals $858.50 but you typed $858.49 — they must match."
    )


def test_row_cap():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(list(config.BOM_TEMPLATE_HEADERS))
    n = config.BOM_MAX_ROWS + 1
    for i in range(n):
        ws.append([f"Part {i}", "", f"P{i}", "V", 1.0, 1, None, ""])
    ws.append([None, None, None, None, None, "Total"])
    items, problems = bom.read_bom_template(_to_bytes(wb))
    assert problems == [f"This BOM has {n} rows — the most is {config.BOM_MAX_ROWS}."]
