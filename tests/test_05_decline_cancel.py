"""Tests for ticket 05: Decline, Cancel, and the processed rename.

Acceptance-criteria coverage:
- Decline writes no Excel row and posts no alert.
- blank_row empties every writable cell; columns A and Z are untouched.
- Cancel after 'processed' (or later) produces the refusal message and leaves the row untouched.
- A buyer's cancel click is refused by the listener.
- A multi-EPIF cancel blanks every row, not just the first.
"""
import json
import os
import shutil
import sys
from unittest.mock import MagicMock, patch

import openpyxl
import pytest

_CURRENT_DIR = os.path.dirname(os.path.abspath(__file__))
PROJECT_ROOT = os.path.dirname(_CURRENT_DIR) if os.path.basename(_CURRENT_DIR) == "tests" else _CURRENT_DIR
SRC_DIR = os.path.join(PROJECT_ROOT, "src")
for p in (PROJECT_ROOT, SRC_DIR):
    if p not in sys.path:
        sys.path.insert(0, p)

try:
    from src import admin, app, lifecycle, log_writer
except ImportError:
    import admin
    import app
    import lifecycle
    import log_writer

SAMPLES = os.path.join(PROJECT_ROOT, "samples")
WORKBOOK = os.path.join(SAMPLES, "Purchasing-Log.xlsx")

FILLED_PDF = os.path.join(SAMPLES, "Prusa_EPIF__2799_PG000025831.pdf")


@pytest.fixture
def workbook(tmp_path):
    if not os.path.exists(WORKBOOK):
        pytest.skip(f"Sample workbook not found at {WORKBOOK}")
    copy = tmp_path / "Purchasing-Log.xlsx"
    shutil.copy(WORKBOOK, copy)
    return str(copy)


@pytest.fixture
def filled():
    if not os.path.exists(FILLED_PDF):
        pytest.skip(f"Sample PDF not found at {FILLED_PDF}")
    from src import epif_parser
    with open(FILLED_PDF, "rb") as fh:
        return epif_parser.parse_epif(fh.read())


# ---------------------------------------------------------------------------
# Decline
# ---------------------------------------------------------------------------

def test_decline_writes_no_excel_row_and_posts_no_alert():
    """handle_decline must update the message but never touch the workbook or DM anyone."""
    client = MagicMock()
    req_data = {
        "parsed": {"item_description": "Resistors", "total_price": 10.0},
        "requester": "Isaac",
        "user_id": "U_CHARLIE",
    }
    history = []

    with patch.object(log_writer, "append_row") as mock_append, \
         patch.object(log_writer, "blank_row") as mock_blank, \
         patch.object(log_writer, "update_row") as mock_update:
        lifecycle.handle_decline(
            client=client,
            channel="C_PURCHASE",
            msg_ts="123.456",
            user_id="U_CHARLIE",
            req_data=req_data,
            history=history,
        )
        mock_append.assert_not_called()
        mock_blank.assert_not_called()
        mock_update.assert_not_called()

    # The message must have been updated (to show "declined")
    client.chat_update.assert_called_once()
    update_kwargs = client.chat_update.call_args[1]
    assert update_kwargs["channel"] == "C_PURCHASE"
    assert update_kwargs["ts"] == "123.456"

    # No DM, no alert channel post
    client.chat_postMessage.assert_not_called()

    # The history now contains a Declined line
    assert any("Declined" in h for h in history)


def test_decline_removes_buttons():
    """The declined-state message must carry no action buttons."""
    client = MagicMock()
    req_data = {"parsed": {"item_description": "Gloves", "total_price": 5.0}, "requester": "Isaac"}

    lifecycle.handle_decline(
        client=client,
        channel="C_PURCHASE",
        msg_ts="1.0",
        user_id="U_CHARLIE",
        req_data=req_data,
        history=[],
    )

    update_kwargs = client.chat_update.call_args[1]
    action_blocks = [b for b in update_kwargs["blocks"] if b.get("type") == "actions"]
    assert action_blocks == [], "declined message must have no action buttons"


def test_decline_restricted_to_approvers():
    """req_decline click from a non-approver is refused; message is not updated."""
    ack = MagicMock()
    respond = MagicMock()
    client = MagicMock()

    body = {
        "user": {"id": "U_BUYER"},
        "channel": {"id": "C_PURCHASE"},
        "message": {"ts": "1.0"},
        "actions": [{"value": json.dumps({"request": {}, "history": []})}],
    }

    with patch.object(admin, "is_approved_reviewer", return_value=False):
        app.handle_req_decline_action(ack, body, respond, client)

    ack.assert_called_once()
    respond.assert_called_once()
    client.chat_update.assert_not_called()


# ---------------------------------------------------------------------------
# Cancel — blank_row workbook behaviour
# ---------------------------------------------------------------------------

def test_blank_row_writable_cells_are_empty_and_columns_a_z_unchanged(workbook, filled):
    """blank_row must clear B–Y and leave A and Z (formulas) intact."""
    row = log_writer.append_row(log_writer.build_row(filled, "Isaac"), workbook_path=workbook)

    # Verify requester written before blanking
    sheet_before = openpyxl.load_workbook(workbook)["Order Log"]
    assert sheet_before[f"B{row}"].value == "Isaac"

    log_writer.blank_row(row, workbook_path=workbook)

    sheet = openpyxl.load_workbook(workbook)["Order Log"]

    # Every writable column (B–Y) must be empty
    for col in "BCDEFGHIJKLMNOPQRSTUVWXY":
        cell_val = sheet[f"{col}{row}"].value
        assert cell_val in (None, ""), (
            f"Column {col} row {row} should be empty after blank_row, got {cell_val!r}"
        )

    # Column A (auto-ID formula) and Z (status formula) must survive
    assert sheet[f"A{row}"].value is not None, "Column A formula must not be erased"
    assert sheet[f"Z{row}"].value is not None, "Column Z formula must not be erased"


# ---------------------------------------------------------------------------
# Cancel — state guard
# ---------------------------------------------------------------------------

def _make_cancel_body(state: str, user_id: str = "U_CHARLIE", thread_ts: str = "100.0"):
    val = json.dumps({
        "state": state,
        "request": {"parsed": {"item_description": "Bolts"}, "requester": "Isaac"},
        "history": [],
    })
    return {
        "user": {"id": user_id},
        "channel": {"id": "C_PURCHASE"},
        "message": {"ts": "100.10"},
        "container": {"thread_ts": thread_ts},
        "actions": [{"value": val}],
    }


def test_cancel_after_processed_produces_refusal_message_and_row_is_untouched(workbook, filled):
    """Cancel on a processed request must post refusal text; the row must not be blanked."""
    row = log_writer.append_row(log_writer.build_row(filled, "Isaac"), workbook_path=workbook)

    client = MagicMock()
    # Simulate thread containing "Logged to row N"
    client.conversations_replies.return_value = {
        "messages": [{"text": f"Logged to row {row}: item — $10.00"}]
    }

    posted_messages = []

    def say(text, **kw):
        posted_messages.append(text)

    with patch.object(log_writer, "blank_row") as mock_blank:
        result = lifecycle.handle_cancel(
            client=client,
            say=say,
            channel="C_PURCHASE",
            thread_ts="100.0",
            msg_ts="100.10",
            user_id="U_CHARLIE",
            req_data={},
            state="processed",
            history=[],
        )
        mock_blank.assert_not_called()

    assert result is False
    assert len(posted_messages) == 1
    assert "purchasing team" in posted_messages[0].lower() or "cannot be cancelled" in posted_messages[0].lower()
    client.chat_update.assert_not_called()

    # Row must be untouched
    sheet = openpyxl.load_workbook(workbook)["Order Log"]
    assert sheet[f"B{row}"].value == "Isaac"


def test_cancel_also_refused_for_confirmed_and_delivered():
    """Cancel is refused for confirmed and delivered states too."""
    for bad_state in ("confirmed", "delivered"):
        client = MagicMock()
        posted = []

        with patch.object(log_writer, "blank_row") as mock_blank:
            result = lifecycle.handle_cancel(
                client=client,
                say=lambda text, **kw: posted.append(text),
                channel="C_PURCHASE",
                thread_ts="100.0",
                msg_ts="100.10",
                user_id="U_CHARLIE",
                req_data={},
                state=bad_state,
                history=[],
            )
            mock_blank.assert_not_called()

        assert result is False, f"cancel should be refused for state={bad_state}"
        assert posted, f"refusal message must be posted for state={bad_state}"


def test_buyer_cancel_click_is_refused():
    """req_cancel click from a buyer (not approver or admin) must be refused."""
    ack = MagicMock()
    respond = MagicMock()
    client = MagicMock()

    body = _make_cancel_body("approved", user_id="U_BUYER")

    with patch.object(admin, "is_approved_reviewer", return_value=False), \
         patch.object(admin, "is_admin_user", return_value=False):
        app.handle_req_cancel_action(ack, body, respond, client)

    ack.assert_called_once()
    respond.assert_called_once()
    # No workbook write, no message update
    client.chat_update.assert_not_called()


# ---------------------------------------------------------------------------
# Cancel — batch (multi-EPIF)
# ---------------------------------------------------------------------------

def test_multi_epif_cancel_blanks_every_row(workbook, filled):
    """A cancel on a batch request must blank all rows written, not just the first."""
    row1 = log_writer.append_row(log_writer.build_row(filled, "Isaac"), workbook_path=workbook)
    row2 = log_writer.append_row(log_writer.build_row(filled, "Isaac"), workbook_path=workbook)

    client = MagicMock()
    # Thread mentions both rows
    client.conversations_replies.return_value = {
        "messages": [
            {"text": f"Logged to row {row1}: EPIF 1 — $10.00"},
            {"text": f"Logged to row {row2}: EPIF 2 — $10.00"},
        ]
    }

    blanked = []

    def fake_blank(row, workbook_path=None):
        blanked.append(row)
        return log_writer.blank_row.__wrapped__(row, workbook_path=workbook) if hasattr(log_writer.blank_row, "__wrapped__") else None

    with patch.object(log_writer, "blank_row", side_effect=lambda r, **kw: blanked.append(r)):
        with patch("src.queue_worker.submit_write_task", side_effect=lambda action_fn, **kw: action_fn()):
            lifecycle.handle_cancel(
                client=client,
                say=lambda text, **kw: None,
                channel="C_PURCHASE",
                thread_ts="100.0",
                msg_ts="100.10",
                user_id="U_CHARLIE",
                req_data={},
                state="approved",
                history=[],
            )

    assert set(blanked) == {row1, row2}, (
        f"Both rows must be blanked; got {blanked}"
    )
